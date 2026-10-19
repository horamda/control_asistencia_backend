"""Jornada laboral, etapa B: motor de calculo de horas desde los fichajes (solo lectura) y su panel.

MySQL real (SH_TEST_DB_PORT), como el resto de los modulos con base. Octubre de 2026: el lunes es el 5."""
import datetime as dt
from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from services import jornada_calculo_service as calc
from services import jornada_service as reglas
from services.asistencia_monthly_report_service import _worked_day
from tests.test_jornada_laboral import QUIEN, _convenio, _form, jl  # noqa: F401  (fixture)
from tests.test_seguridad import database  # noqa: F401  (fixture)

D = dt.date
HOY = dt.datetime(2026, 10, 25, 10, 0)
DESDE, HASTA = D(2026, 10, 5), D(2026, 10, 20)


@pytest.fixture
def hj(jl, monkeypatch):  # noqa: F811
    """Empleado 10 con horario L-V 08-12 y 13-17 (del fixture jl); 11, 12 y 13 sin horario: les vale la regla de la
    empresa (L-V, 8 h). El 12 tiene una condicion individual de 6 h desde el 8/10 y el 13 se dio de baja el 7/10."""
    conn = jl()
    c = conn.cursor()
    for sql in [
        "ALTER TABLE empleados ADD COLUMN fecha_baja DATE NULL",
        "ALTER TABLE empleados ADD COLUMN reporta_a_empleado_id INT NULL",
        "UPDATE empleados SET activo = 0, fecha_baja = '2026-10-07' WHERE id = 13",
        "CREATE TABLE asistencia_marcas (id BIGINT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, empleado_id INT, fecha DATE,"
        " hora TIME, accion VARCHAR(10), metodo VARCHAR(20), tipo_marca VARCHAR(20) DEFAULT 'jornada')",
        "CREATE TABLE empleado_excepciones (id INT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, empleado_id INT, fecha DATE,"
        " tipo VARCHAR(30), descripcion VARCHAR(200) NULL, anula_horario TINYINT DEFAULT 0)",
        "CREATE TABLE excepcion_bloques (id INT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, excepcion_id INT, orden INT,"
        " hora_entrada TIME, hora_salida TIME)",
        "CREATE TABLE asistencia_dias_no_laborables (id INT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, sucursal_id INT,"
        " sector_id INT DEFAULT 0, fecha DATE, motivo VARCHAR(120))",
        "CREATE TABLE justificaciones (id INT AUTO_INCREMENT PRIMARY KEY, empleado_id INT, fecha DATE, fecha_desde DATE NULL,"
        " fecha_hasta DATE NULL, estado VARCHAR(20))",
        "CREATE TABLE IF NOT EXISTS vacaciones (id INT AUTO_INCREMENT PRIMARY KEY, empleado_id INT, fecha_desde DATE, fecha_hasta DATE)",
        "CREATE TABLE IF NOT EXISTS francos (id INT AUTO_INCREMENT PRIMARY KEY, empleado_id INT, fecha DATE)",
        "INSERT INTO asistencia_dias_no_laborables (empresa_id, sucursal_id, fecha, motivo) VALUES"
        " (0, 0, '2026-10-12', 'Diversidad cultural'), (1, 2, '2026-10-13', 'Feriado local')",
        "INSERT INTO justificaciones (empleado_id, fecha, estado) VALUES (10, '2026-10-14', 'aprobada'), (10, '2026-10-13', 'pendiente')",
        "INSERT INTO vacaciones (empleado_id, fecha_desde, fecha_hasta) VALUES (10, '2026-10-15', '2026-10-15')",
        "INSERT INTO francos (empleado_id, fecha) VALUES (10, '2026-10-16')",
        "INSERT INTO empleado_excepciones (id, empresa_id, empleado_id, fecha, tipo, anula_horario) VALUES"
        " (1, 1, 10, '2026-10-19', 'CAMBIO_HORARIO', 0), (2, 1, 10, '2026-10-20', 'LICENCIA', 1)",
        "INSERT INTO excepcion_bloques (empresa_id, excepcion_id, orden, hora_entrada, hora_salida) VALUES (1, 1, 1, '10:00', '14:00')",
    ]:
        c.execute(sql)
    from scripts.migrate_20261019_02_jornada_calculo import migrate_cursor
    migrate_cursor(c)
    migrate_cursor(c)  # idempotente
    from scripts.migrate_20261019_04_banco_horas import migrate_cursor as etapa_d
    etapa_d(c)
    etapa_d(c)  # idempotente
    conn.commit()
    c.close()
    conn.close()
    monkeypatch.setattr(calc, "now_local", lambda: HOY)
    reglas.guardar_regla(QUIEN, "empresa", None, _form(vigente_desde="2026-01-01", horas_diarias="8", horas_semanales="48",
                                                       dias_laborales=["1", "2", "3", "4", "5"]))
    reglas.guardar_regla(QUIEN, "empleado", 12, _form(vigente_desde="2026-10-08", horas_diarias="6"))
    return jl


def marcar(empleado_id, *marcas, empresa_id=1):
    """marcas: ("2026-10-05", "08:00", "ingreso"[, metodo])."""
    with calc.db.transaction() as c:
        for m in marcas:
            fecha, hora, accion = m[:3]
            c.execute("INSERT INTO asistencia_marcas (empresa_id, empleado_id, fecha, hora, accion, metodo) VALUES (%s,%s,%s,%s,%s,%s)",
                      (empresa_id, empleado_id, fecha, hora, accion, m[3] if len(m) > 3 else "qr"))


def jornada(empleado_id, fecha, *tramos):
    """Pares ingreso/egreso del mismo dia: jornada(10, "2026-10-05", ("08:00", "12:00"), ...)."""
    marcar(empleado_id, *[(fecha, h, a) for t in tramos for h, a in zip(t, ("ingreso", "egreso"))])


def dias(datos, empleado_id):
    fila = next(r for r in datos["empleados"] if r["empleado"]["id"] == empleado_id)
    return {d["fecha"].day: d for d in fila["dias"]}, fila


def _cargar_empleado_10():
    jornada(10, "2026-10-05", ("08:00", "12:00"), ("13:00", "17:00"))
    jornada(10, "2026-10-06", ("08:00", "12:00"), ("13:00", "19:00"))
    jornada(10, "2026-10-07", ("08:05", "12:00"), ("13:00", "17:00"))
    marcar(10, ("2026-10-08", "08:00", "ingreso"))
    marcar(10, ("2026-10-09", "08:00", "ingreso"), ("2026-10-09", "17:00", "egreso", "automatico"))
    jornada(10, "2026-10-12", ("08:00", "12:00"))
    jornada(10, "2026-10-17", ("09:00", "11:00"))
    jornada(10, "2026-10-19", ("10:00", "14:00"))
    jornada(10, "2026-10-20", ("08:00", "12:00"))


# --- Emparejado (sin base) ---------------------------------------------------------------------------------------

def test_emparejar_cruza_medianoche_y_respeta_el_maximo_de_turno():
    m = [dict(id=1, fecha=D(2026, 10, 5), hora="22:00", accion="ingreso"), dict(id=2, fecha=D(2026, 10, 6), hora="06:00", accion="egreso"),
         dict(id=3, fecha=D(2026, 10, 7), hora="08:00", accion="ingreso"), dict(id=4, fecha=D(2026, 10, 8), hora="09:00", accion="egreso"),
         dict(id=5, fecha=D(2026, 10, 9), hora="17:00", accion="egreso")]
    pares, sueltas = calc.emparejar(m)
    assert [(i.isoformat(), f.isoformat()) for i, f, _ in pares[D(2026, 10, 5)]] == [("2026-10-05T22:00:00", "2026-10-06T06:00:00")]
    assert D(2026, 10, 6) not in pares  # el turno cuenta para el dia del ingreso
    assert [len(sueltas[D(2026, 10, d)]) for d in (7, 8, 9)] == [1, 1, 1]  # 25 h no es un turno: quedan sin pareja


def test_horas_formato_24h():
    assert (calc.horas(0), calc.horas(475), calc.horas(-90), calc.horas(1500)) == ("0:00", "7:55", "-1:30", "25:00")


# --- Motor --------------------------------------------------------------------------------------------------------

def test_dia_normal_excedente_tolerancia_e_incompletos(hj):
    _cargar_empleado_10()
    d, fila = dias(calc.calcular(1, DESDE, HASTA), 10)
    lunes = d[5]
    assert (lunes["estado"], lunes["programados"], lunes["trabajados"], lunes["excedente"], lunes["faltantes"], lunes["descanso"]) == \
        ("normal", 480, 480, 0, 0, 60)
    assert lunes["bloques"] == ["08:00–12:00", "13:00–17:00"] and lunes["horario"] == "Administrativo"
    assert (d[6]["trabajados"], d[6]["excedente"], d[6]["ordinarios"]) == (600, 120, 480)  # excedente, no hora extra
    # La tolerancia de 10 min del horario no cambia los minutos: llego 5 min tarde y faltan 5.
    assert (d[7]["estado"], d[7]["trabajados"], d[7]["faltantes"]) == ("normal", 475, 5)
    assert (d[8]["estado"], d[8]["excedente"], d[8]["faltantes"]) == ("revisar", 0, 0) and "sin pareja" in d[8]["motivo"]
    assert d[9]["estado"] == "revisar" and "egreso automático" in d[9]["motivo"] and d[9]["pares"] == ["08:00–17:00 auto"]
    assert d[10]["estado"] == d[11]["estado"] == "descanso"
    assert any(a["tipo"] == "exceso_diario" and a["fecha"].day == 6 for a in fila["alertas"])


def test_feriados_licencias_ausencias_y_excepciones(hj):
    _cargar_empleado_10()
    d, fila = dias(calc.calcular(1, DESDE, HASTA), 10)
    assert (d[12]["tipo_dia"], d[12]["programados"], d[12]["excedente"], d[12]["feriado"]) == ("feriado", 0, 240, "Diversidad cultural")
    assert (d[13]["estado"], d[13]["faltantes"]) == ("ausente", 480)  # la justificacion pendiente no cuenta
    assert [(d[x]["estado"], d[x]["licencia"], d[x]["faltantes"]) for x in (14, 15, 16)] == [
        ("licencia", "Justificación", 0), ("licencia", "Vacaciones", 0), ("licencia", "Franco", 0)]
    assert (d[17]["tipo_dia"], d[17]["excedente"]) == ("no_laborable", 120)  # sabado sin horario
    assert (d[19]["programados"], d[19]["bloques"], d[19]["estado"]) == (240, ["10:00–14:00"], "normal")  # CAMBIO_HORARIO
    assert (d[20]["tipo_dia"], d[20]["estado"]) == ("licencia", "revisar") and "licencia" in d[20]["motivo"]
    t = fila["totales"]
    assert (t["dias_revisar"], t["dias_ausente"], t["dias_licencia"], t["feriados_trabajados"]) == (3, 1, 3, 1)


def test_turno_nocturno_que_cruza_la_medianoche_y_descanso_corto(hj):
    marcar(11, ("2026-10-05", "22:00", "ingreso"), ("2026-10-06", "06:00", "egreso"))
    jornada(11, "2026-10-06", ("14:00", "18:00"))
    marcar(11, ("2026-10-07", "08:00", "ingreso"), ("2026-10-08", "09:00", "egreso"))
    d, fila = dias(calc.calcular(1, DESDE, HASTA), 11)
    noche = d[5]
    assert (noche["estado"], noche["trabajados"], noche["programados"], noche["cruza_medianoche"]) == ("normal", 480, 480, True)
    assert noche["nocturnos"] == 480 and noche["pares"] == ["22:00–06:00 (+1)"]  # franja de la LCT: 21 a 6
    assert (d[6]["trabajados"], d[6]["faltantes"]) == (240, 240)
    assert any(a["tipo"] == "descanso_corto" and a["fecha"].day == 6 for a in fila["alertas"])
    assert d[7]["estado"] == d[8]["estado"] == "revisar"


def test_cambio_de_condicion_a_mitad_del_periodo_y_baja(hj):
    d, _ = dias(calc.calcular(1, DESDE, HASTA), 12)
    assert [d[x]["programados"] for x in (5, 6, 7, 8, 9)] == [480, 480, 480, 360, 360]
    assert d[7]["reglas"] != d[8]["reglas"]  # cada dia guarda las versiones de regla que uso
    assert d[13]["tipo_dia"] == "feriado"  # feriado de su sucursal
    baja, _ = dias(calc.calcular(1, DESDE, HASTA), 13)
    assert (baja[7]["estado"], baja[8]["estado"], baja[8]["programados"]) == ("ausente", "fuera", 0)


def test_filtros_por_sucursal_convenio_y_aislamiento(hj):
    camioneros = _convenio("CCT_40_89")["id"]
    reglas.asignar_convenio(QUIEN, [10], camioneros, "2026-01-01")
    assert [r["empleado"]["id"] for r in calc.calcular(1, DESDE, HASTA, {"sucursal_id": "2"})["empleados"]] == [12]
    assert [r["empleado"]["id"] for r in calc.calcular(1, DESDE, HASTA, {"convenio_id": str(camioneros)})["empleados"]] == [10]
    marcar(20, ("2026-10-05", "08:00", "ingreso"), ("2026-10-05", "16:00", "egreso"), empresa_id=2)
    assert 20 not in [r["empleado"]["id"] for r in calc.calcular(1, DESDE, HASTA)["empleados"]]
    with pytest.raises(calc.Error, match="máximo"):
        calc.calcular(1, D(2026, 1, 1), D(2026, 6, 30))


def test_coincide_con_el_reporte_mensual_en_jornadas_del_mismo_dia(hj):
    _cargar_empleado_10()
    d, _ = dias(calc.calcular(1, DESDE, HASTA), 10)
    with calc.db.transaction(read_only=True) as c:
        for dia in (5, 6, 7, 12, 17):
            marcas = calc.db.all_rows(c, "SELECT id, hora, accion FROM asistencia_marcas WHERE empleado_id = 10 AND fecha = %s",
                                      (D(2026, 10, dia),))
            assert _worked_day(marcas)[0] == round(d[dia]["fichados"] / 60, 2), dia


def test_registrar_conserva_historial_y_detecta_fichajes_cambiados(hj):
    _cargar_empleado_10()
    with calc.db.transaction(read_only=True) as c:
        antes = calc.db.all_rows(c, "SELECT * FROM asistencia_marcas ORDER BY id")
    primero, datos = calc.registrar(QUIEN, DESDE, HASTA, motivo="Cierre")
    registrado, huellas = calc.ultimo_calculo(1, DESDE, HASTA)
    assert registrado["id"] == primero and registrado["motivo"] == "Cierre" and registrado["empleados"] == 4
    assert all(huellas[(r["empleado"]["id"], d["fecha"])] == d["huella"] for r in datos["empleados"] for d in r["dias"])
    marcar(10, ("2026-10-08", "17:00", "egreso"))  # alguien completa la fichada que faltaba
    nuevo = calc.calcular(1, DESDE, HASTA)
    cambiados = [(r["empleado"]["id"], d["fecha"].day) for r in nuevo["empleados"] for d in r["dias"]
                 if huellas[(r["empleado"]["id"], d["fecha"])] != d["huella"]]
    assert cambiados == [(10, 7), (10, 8)]  # el dia y el anterior (por un posible turno de medianoche)
    segundo, _ = calc.registrar(QUIEN, DESDE, HASTA)
    with calc.db.transaction(read_only=True) as c:
        filas = calc.db.all_rows(c, """SELECT calculo_id, estado, minutos_trabajados FROM jl_resultados_dia
                                       WHERE empleado_id = 10 AND fecha = '2026-10-08' ORDER BY calculo_id""")
        assert calc.db.all_rows(c, "SELECT * FROM asistencia_marcas WHERE id <= %s ORDER BY id", (antes[-1]["id"],)) == antes
    assert [(f["calculo_id"], f["estado"], f["minutos_trabajados"]) for f in filas] == [(primero, "revisar", 0), (segundo, "normal", 540)]
    assert [k["id"] for k in calc.calculos(1)] == [segundo, primero]


# --- Panel ------------------------------------------------------------------------------------------------------

@pytest.fixture
def client(hj, monkeypatch):
    import web.auth.decorators as auth
    from web.horas import routes

    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.secret_key = "test"
    app.config["TESTING"] = True
    app.register_blueprint(routes.web)
    permisos = {"acciones": {"ver", "editar", "exportar"}}

    def can_web(code, action="ver"):
        return code == "horas_jornada" and action in permisos["acciones"]

    app.jinja_env.globals.update(csp_nonce=lambda: "n", csrf_token=lambda: "t", can_web=can_web,
                                 export_toolbar={"enabled": False, "items": []})
    monkeypatch.setattr(auth, "can_access_module", lambda uid, module, action="ver": can_web(module, action))
    monkeypatch.setattr(routes, "_cached_web_user", lambda uid: dict(id=99, empresa_id=1, activo=1))
    monkeypatch.setattr(routes, "log_audit", lambda *a, **kw: None)
    test_client = app.test_client()
    with test_client.session_transaction() as sess:
        sess["user_id"] = 99
    test_client.permisos = permisos
    return test_client


def test_pantallas_registro_exportacion_y_permisos(client):
    _cargar_empleado_10()
    base = "/horas-jornada"
    pagina = client.get(base + "/?mes=2026-10").get_data(as_text=True)
    assert "Horas trabajadas" in pagina and "Perez, Ana" in pagina and "no tiene cálculo registrado" in pagina
    assert "Registrar cálculo" in pagina and "Exportar Excel" in pagina
    detalle = client.get(base + "/empleado/10?mes=2026-10")
    texto = detalle.get_data(as_text=True)
    assert detalle.status_code == 200 and "Lun 05/10" in texto and "Semana del 05/10" in texto and "08:00–17:00 auto" in texto
    assert client.get(base + "/empleado/20?mes=2026-10").status_code == 302  # otra empresa
    r = client.post(base + "/registrar", data={"mes": "2026-10", "motivo": "Cierre"})
    assert r.status_code == 302 and calc.ultimo_calculo(1, D(2026, 10, 1), D(2026, 10, 31))[0]["motivo"] == "Cierre"
    assert "no cambiaron desde ese registro" in client.get(base + "/?mes=2026-10").get_data(as_text=True)
    xlsx = client.get(base + "/exportar?mes=2026-10&sucursal_id=1")
    libro = load_workbook(BytesIO(xlsx.data))
    assert xlsx.status_code == 200 and libro.sheetnames == ["Resumen", "Detalle diario"]
    resumen = {fila[1]: fila for fila in libro["Resumen"].iter_rows(min_row=5, values_only=True)}
    assert set(resumen) == {"Perez, Ana", "Garcia, Luis", "Ajena, Persona"}  # sucursal 1, con la baja del mes
    assert resumen["Perez, Ana"][5] == calc.horas(calc.calcular(1, D(2026, 10, 1), D(2026, 10, 31))["empleados"][2]["totales"]["trabajados"])
    client.permisos["acciones"] = {"ver"}
    pagina = client.get(base + "/?mes=2026-10").get_data(as_text=True)
    assert "Registrar cálculo" not in pagina and "Exportar Excel" not in pagina
    assert client.post(base + "/registrar", data={"mes": "2026-10"}).status_code == 403
    assert client.get(base + "/exportar?mes=2026-10").status_code == 403
    client.permisos["acciones"] = set()
    assert client.get(base + "/").status_code == 403
