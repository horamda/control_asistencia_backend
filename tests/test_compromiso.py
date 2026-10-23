"""Indice de compromiso: carga (a mano y planilla), totales por promedio, lo que ve cada empleado y el panel."""
import io
from contextlib import contextmanager
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import Workbook, load_workbook

import app as app_module
import routes.mobile_v1_routes as mobile_routes
import utils.jwt_guard as jwt_guard
import web.auth.decorators as auth
import web.carga.carga_routes as carga_web
from services import compromiso_service as cmp
from tests.test_seguridad import database  # noqa: F401  (fixture)


@pytest.fixture
def base(database, monkeypatch):  # noqa: F811
    from scripts.migrate_20261023_01_compromiso import migrate_cursor

    conn = database()
    c = conn.cursor()
    for sql in [
        "ALTER TABLE sucursales ADD COLUMN empresa_id INT NOT NULL DEFAULT 1",
        "ALTER TABLE sucursales ADD COLUMN activa TINYINT NOT NULL DEFAULT 1",
        "UPDATE sucursales SET empresa_id = 2 WHERE id = 3",
        "ALTER TABLE sectores ADD COLUMN empresa_id INT NOT NULL DEFAULT 1",
        "ALTER TABLE sectores ADD COLUMN activo TINYINT NOT NULL DEFAULT 1",
        "INSERT INTO sectores (id, nombre) VALUES (2, 'Administración')",
        "ALTER TABLE empresas ADD COLUMN razon_social VARCHAR(100) NULL",
        "UPDATE empresas SET razon_social = 'Distribuidora Sur' WHERE id = 1",
    ]:
        c.execute(sql)
    migrate_cursor(c)
    conn.commit()
    c.close()
    conn.close()

    @contextmanager
    def transaction(*, read_only=False):
        conn = database()
        cursor = conn.cursor(dictionary=True)
        try:
            yield cursor
            if not read_only:
                conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cursor.close()
            conn.close()

    monkeypatch.setattr(cmp.db, "transaction", transaction)
    return database


def empleado(eid=10, empresa=1, sucursal=1, sector=1):
    return dict(id=eid, empresa_id=empresa, sucursal_id=sucursal, sector_id=sector, activo=1)


def test_guardar_y_totales_por_promedio(base):
    cmp.guardar(1, 99, 2026, 1, 1, "70")
    cmp.guardar(1, 99, 2026, 1, 2, 80)
    cmp.guardar(1, 99, 2026, 2, 1, "60,5")
    datos = cmp.listado(1, 2026)
    central = next(s for s in datos["sucursales"] if s["sucursal_id"] == 1)
    assert central["indice"] == 75.0 and central["calculado"] is True
    assert datos["empresa"] == dict(indice=70.2, calculado=True)  # promedio de las 3 areas
    # Un total cargado manda sobre el promedio; volver a guardar reemplaza.
    cmp.guardar(1, 99, 2026, 1, 0, 90)
    cmp.guardar(1, 99, 2026, 0, 0, 88)
    cmp.guardar(1, 99, 2026, 1, 1, 72)
    datos = cmp.listado(1, 2026)
    central = next(s for s in datos["sucursales"] if s["sucursal_id"] == 1)
    assert central["indice"] == 90.0 and central["calculado"] is False
    assert datos["empresa"] == dict(indice=88.0, calculado=False)
    assert len(datos["filas"]) == 5 and cmp.anios(1) == [2026]


def test_validaciones(base):
    with pytest.raises(cmp.CompromisoError, match="no es de esta empresa"):
        cmp.guardar(1, 99, 2026, 3, 0, 50)  # sucursal de la empresa 2
    with pytest.raises(cmp.CompromisoError, match="elegí también la sucursal"):
        cmp.guardar(1, 99, 2026, 0, 1, 50)
    for malo in ("150", "-1", "abc"):
        with pytest.raises(cmp.CompromisoError):
            cmp.guardar(1, 99, 2026, 1, 1, malo)
    with pytest.raises(cmp.CompromisoError, match="Año"):
        cmp.guardar(1, 99, 1999, 1, 1, 50)
    assert cmp._indice("0,72") == 72.0 and cmp._indice("72%") == 72.0 and cmp._indice("1") == 1.0


def _xlsx(filas):
    libro = Workbook()
    hoja = libro.active
    for fila in filas:
        hoja.append(fila)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def test_importar_planilla_y_csv(base):
    contenido = _xlsx([["Año", "Sucursal", "Área", "Índice"],
                       [2025, "central", "ENTREGA", 0.65],
                       [2025, "Central", "administracion", "70%"],
                       [2025, "Otra", "Entrega", 55],
                       [2025, "Central", "", 68],  # total de la sucursal
                       [2025, None, None, 61],      # total de la empresa
                       [None, None, None, None]])
    assert cmp.importar(1, 99, "datos.xlsx", contenido) == dict(filas=5, anios=[2025])
    datos = cmp.listado(1, 2025)
    assert {(f["sucursal_id"], f["sector_id"]): float(f["indice"]) for f in datos["filas"]} == {
        (1, 1): 65.0, (1, 2): 70.0, (2, 1): 55.0, (1, 0): 68.0, (0, 0): 61.0}
    csv_texto = "Año;Sucursal;Área;Índice\n2026;Central;Entrega;80\n".encode("utf-8")
    assert cmp.importar(1, 99, "datos.csv", csv_texto)["filas"] == 1


def test_importar_con_errores_no_guarda_nada(base):
    contenido = _xlsx([["Año", "Sucursal", "Área", "Índice"],
                       [2025, "Central", "Entrega", 70],
                       [2025, "Inexistente", "Entrega", 70],
                       [2025, "Central", "Cocina", 70],
                       [2025, "Empresa B", "Entrega", 70],  # sucursal de otra empresa
                       [2025, "Central", "Entrega", 170]])
    with pytest.raises(cmp.CompromisoError) as exc:
        cmp.importar(1, 99, "datos.xlsx", contenido)
    texto = str(exc.value)
    assert "No se importó nada" in texto and "Fila 3" in texto and "Fila 4" in texto and "Fila 5" in texto and "Fila 6" in texto
    assert cmp.listado(1, 2025)["filas"] == []
    with pytest.raises(cmp.CompromisoError, match="Faltan columnas"):
        cmp.importar(1, 99, "x.csv", b"Sucursal;Valor\nCentral;3\n")


def test_lo_que_ve_cada_empleado(base):
    cmp.guardar(1, 99, 2025, 1, 1, 60)
    cmp.guardar(1, 99, 2025, 1, 2, 70)
    cmp.guardar(1, 99, 2026, 1, 1, 72)
    cmp.guardar(1, 99, 2026, 1, 2, 80)
    cmp.guardar(1, 99, 2026, 2, 1, 50)
    vista = cmp.para_empleado(empleado())
    assert vista["anio"] == 2026 and vista["anios_disponibles"] == [2026, 2025]
    assert vista["area"] == dict(nombre="Entrega", indice=72.0, calculado=False, anterior=60.0, variacion=12.0)
    assert vista["sucursal"] == dict(nombre="Central", indice=76.0, calculado=True, anterior=65.0, variacion=11.0)
    assert vista["empresa"]["indice"] == 67.3 and vista["empresa"]["anterior"] == 65.0
    assert vista["empresa"]["nombre"] == "Distribuidora Sur"
    otra = cmp.para_empleado(empleado(12, sucursal=2, sector=2), 2026)  # area sin dato
    assert otra["area"]["indice"] is None and otra["sucursal"]["indice"] == 50.0
    sin_datos = cmp.para_empleado(empleado(20, empresa=2, sucursal=3))
    assert sin_datos["anios_disponibles"] == [] and sin_datos["empresa"]["indice"] is None


def test_plantilla_trae_sucursales_y_areas_con_empleados(base):
    libro = load_workbook(io.BytesIO(cmp.plantilla(1, 2026)))
    filas = list(libro["Compromiso"].iter_rows(values_only=True))
    assert filas[0] == cmp.COLUMNAS
    assert (2026, "Central", "Entrega", None) in filas and (2026, "Otra", "Entrega", None) in filas
    assert all(f[1] != "Empresa B" for f in filas)


def test_panel(base, monkeypatch):
    from web.compromiso.compromiso_routes import compromiso_bp

    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.secret_key = "test"
    app.config["TESTING"] = True
    app.register_blueprint(compromiso_bp)
    app.jinja_env.globals.update(csp_nonce=lambda: "n", csrf_token=lambda: "t", can_web=lambda code, *a: code == "compromiso",
                                 export_toolbar={"enabled": False, "items": []})
    monkeypatch.setattr(auth, "can_access_module", lambda *a: True)
    monkeypatch.setattr(carga_web, "get_user", lambda uid: dict(id=uid, empresa_id=1, rol="rrhh", activo=1))
    monkeypatch.setattr("web.compromiso.compromiso_routes.log_audit", lambda *a, **k: None)
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = 99
    assert "No hay datos cargados" in client.get("/compromiso/").get_data(as_text=True)
    r = client.post("/compromiso/guardar", data=dict(anio=2026, sucursal_id=1, sector_id=1, indice="70"))
    assert r.status_code == 302
    html = client.get("/compromiso/?anio=2026").get_data(as_text=True)
    assert "Resumen 2026" in html and "70.0%" in html and "Promedio de 1 área" in html
    fila = cmp.listado(1, 2026)["filas"][0]
    r = client.post("/compromiso/importar", data=dict(archivo=(io.BytesIO(_xlsx(
        [["Año", "Sucursal", "Área", "Índice"], [2026, "Otra", "Entrega", 40]])), "d.xlsx")),
        content_type="multipart/form-data")
    assert r.status_code == 302 and len(cmp.listado(1, 2026)["filas"]) == 2
    client.post(f"/compromiso/{fila['id']}/eliminar", data=dict(anio=2026))
    assert len(cmp.listado(1, 2026)["filas"]) == 1
    r = client.get("/compromiso/plantilla.xlsx")
    assert r.status_code == 200 and r.mimetype.endswith("spreadsheetml.sheet")
    monkeypatch.setattr(auth, "can_access_module", lambda uid, module, action="ver": action == "ver")
    assert client.post("/compromiso/guardar", data=dict(anio=2026, sucursal_id=1, sector_id=1, indice=1)).status_code == 403


def test_endpoint_de_la_app(monkeypatch):
    monkeypatch.setattr(app_module, "init_db", lambda: None)
    monkeypatch.setattr(mobile_routes, "get_profile_photo_version_by_dni", lambda dni: None)
    monkeypatch.setattr(jwt_guard, "verificar_token", lambda token: {"empleado_id": 10})
    monkeypatch.setattr(mobile_routes, "get_empleado_by_id", lambda _: empleado())
    pedidos = []
    monkeypatch.setattr(cmp, "para_empleado", lambda emp, anio=None: pedidos.append(anio) or {"anio": anio or 2026})
    client = app_module.create_app().test_client()
    headers = {"Authorization": "Bearer abc"}
    assert client.get("/api/v1/mobile/me/compromiso", headers=headers).json == {"anio": 2026}
    assert client.get("/api/v1/mobile/me/compromiso?anio=2025", headers=headers).status_code == 200
    assert client.get("/api/v1/mobile/me/compromiso?anio=19", headers=headers).status_code == 400
    assert pedidos == [None, 2025]
    assert client.get("/api/v1/mobile/me/compromiso").status_code == 401
