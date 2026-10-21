"""Historial de marcas: calidad de la ubicacion (contrato 1.54.0) en la pantalla y el Excel."""
import io
from decimal import Decimal
import re

from openpyxl import load_workbook

import app as app_module
import web.auth.decorators as auth_decorators
import web.asistencias.asistencias_routes as asistencias_routes

BASE = dict(empresa_id=1, empleado_id=2, empresa_nombre="Empresa Test", apellido="Perez", nombre="Ana",
            dni="30111222", fecha="2026-10-20", accion="ingreso", tipo_marca="jornada", metodo="qr", gps_ok=1,
            gps_distancia_m=3.2, gps_tolerancia_m=30.0, lat=-36.3, lon=-57.6, estado="ok", observaciones="",
            fecha_creacion="2026-10-20 08:00:01", corregida_manualmente=0, es_resumen=0)
FILAS = [
    dict(BASE, id=1, hora="08:00:00", gps_precision_m=8.4, gps_simulada=1, gps_edad_s=300),
    dict(BASE, id=2, hora="08:05:00", gps_precision_m=12.0, gps_simulada=0, gps_edad_s=10),
    # Fichada de una version anterior de la app y resumen historico: sin datos de calidad.
    dict(BASE, id=3, hora="08:10:00", gps_precision_m=None, gps_simulada=None, gps_edad_s=None),
    dict(BASE, id=None, hora="08:15:00", es_resumen=1),
]


def _client(monkeypatch):
    monkeypatch.setattr(app_module, "init_db", lambda: None)
    monkeypatch.setattr(asistencias_routes, "get_sectores", lambda include_inactive=True: [])
    monkeypatch.setattr(auth_decorators, "has_role", lambda actor_id, role: True)
    monkeypatch.setattr(asistencias_routes, "get_empleados", lambda include_inactive=True: [])
    monkeypatch.setattr(asistencias_routes, "get_empresas", lambda include_inactive=True: [])
    app = app_module.create_app()
    app.config["TESTING"] = True
    app.config["WTF_CSRF_ENABLED"] = False
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = 99
    return client


def _texto(celda):
    return " ".join(re.sub(r"<[^>]+>", " ", celda).split())


def test_historial_muestra_precision_simulado_y_antiguedad(monkeypatch):
    client = _client(monkeypatch)
    monkeypatch.setattr(asistencias_routes, "get_marcas_admin_page", lambda **kwargs: (FILAS, len(FILAS)))
    html = client.get("/asistencias/marcas").get_data(as_text=True)
    celdas = [_texto(c) for c in re.findall(r'<td data-label="GPS">(.*?)</td>', html, re.S)]
    assert celdas == [
        "ok simulado ±8 m · posición de hace 5 min",
        "ok ±12 m",  # 10 s: no se marca como vieja
        "ok",
        "ok",
    ], celdas


def test_excel_suma_columnas_de_calidad(monkeypatch):
    client = _client(monkeypatch)
    monkeypatch.setattr(asistencias_routes, "get_marcas_admin_export", lambda **kwargs: FILAS)
    resp = client.get("/asistencias/marcas.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data))["Historial"]
    filas = list(ws.iter_rows(values_only=True))
    encabezado = next(r for r in filas if r and r[0] == "ID")
    i = encabezado.index("Precision (m)")
    assert encabezado[i - 2:i + 4] == ("Lat", "Lon", "Precision (m)", "GPS simulado", "Antiguedad GPS (s)", "Estado")
    detalle = filas[filas.index(encabezado) + 1:]
    assert [tuple(r[i:i + 3]) for r in detalle] == [(8.4, "si", 300), (12.0, "no", 10), (None, None, None),
                                                  (None, None, None)]


# --- MySQL real (SH_TEST_DB_PORT): la vista del reporte con el join por id ------------------------------------
from tests.test_seguridad import database  # noqa: E402,F401  (fixture)


def test_consultas_del_historial_traen_la_calidad(database, monkeypatch):  # noqa: F811
    import repositories.asistencia_marca_repository as repo
    from scripts import migrate_20261020_02_vista_reporte_marcas as vista

    conn = database()
    c = conn.cursor()
    lado = lambda s: (f"hora_{s} TIME NULL, metodo_{s} VARCHAR(20) NULL, lat_{s} DECIMAL(10,7) NULL, "
                      f"lon_{s} DECIMAL(10,7) NULL, foto_{s} TEXT NULL, gps_ok_{s} TINYINT NULL, "
                      f"gps_distancia_{s}_m DECIMAL(10,2) NULL, gps_tolerancia_{s}_m DECIMAL(10,2) NULL, "
                      f"gps_ref_lat_{s} DECIMAL(10,7) NULL, gps_ref_lon_{s} DECIMAL(10,7) NULL")
    for sql in [
        "ALTER TABLE empresas ADD COLUMN razon_social VARCHAR(100) NULL",
        "ALTER TABLE empleados ADD COLUMN dni VARCHAR(20) NULL",
        """CREATE TABLE asistencia_marcas (id BIGINT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, empleado_id INT,
             asistencia_id INT NULL, fecha DATE, hora TIME, accion VARCHAR(10), metodo VARCHAR(20), tipo_marca VARCHAR(20),
             lat DECIMAL(10,7) NULL, lon DECIMAL(10,7) NULL, foto TEXT NULL, gps_ok TINYINT NULL,
             gps_distancia_m DECIMAL(10,2) NULL, gps_tolerancia_m DECIMAL(10,2) NULL, gps_ref_lat DECIMAL(10,7) NULL,
             gps_ref_lon DECIMAL(10,7) NULL, estado VARCHAR(30) NULL, observaciones TEXT NULL, fecha_creacion DATETIME NULL,
             corregida_manualmente TINYINT DEFAULT 0, gps_precision_m DECIMAL(8,1) NULL, gps_simulada TINYINT(1) NULL,
             gps_edad_s INT NULL)""",
        f"""CREATE TABLE asistencias (id INT AUTO_INCREMENT PRIMARY KEY, empresa_id INT, empleado_id INT, fecha DATE,
             {lado('entrada')}, {lado('salida')}, estado VARCHAR(30) NULL, observaciones TEXT NULL,
             created_at DATETIME NULL)""",
        """INSERT INTO asistencia_marcas (empresa_id, empleado_id, fecha, hora, accion, metodo, tipo_marca, lat, lon,
             gps_ok, gps_precision_m, gps_simulada, gps_edad_s) VALUES
             (1, 10, '2026-10-20', '08:00', 'ingreso', 'qr', 'jornada', -36.3, -57.6, 1, 8.4, 1, 300),
             (1, 10, '2026-10-20', '17:00', 'egreso', 'qr', 'jornada', -36.3, -57.6, 1, NULL, NULL, NULL)""",
        # Dia sin marcas: la vista lo arma desde el resumen (id NULL, sin calidad).
        """INSERT INTO asistencias (empresa_id, empleado_id, fecha, hora_entrada, metodo_entrada)
             VALUES (1, 11, '2026-10-19', '09:00', 'manual')""",
    ]:
        c.execute(sql)
    vista.migrate_cursor(c)  # base sin la vista: el bootstrap la crea
    conn.commit()
    c.close()
    conn.close()
    monkeypatch.setattr(repo, "get_db", database)

    rows, total = repo.get_page_admin(page=1, per_page=20)
    assert total == 3
    calidad = {(str(r["fecha"]), str(r["hora"])): (r["gps_precision_m"], r["gps_simulada"], r["gps_edad_s"]) for r in rows}
    assert calidad == {("2026-10-20", "8:00:00"): (Decimal("8.4"), 1, 300), ("2026-10-20", "17:00:00"): (None, None, None),
                       ("2026-10-19", "9:00:00"): (None, None, None)}
    exportadas = repo.get_for_export_admin()
    assert len(exportadas) == 3
    assert {(str(r["hora"]), r["gps_edad_s"]) for r in exportadas} == {("8:00:00", 300), ("17:00:00", None),
                                                                        ("9:00:00", None)}


def test_bootstrap_crea_solo_lo_que_falta_y_no_toca_una_vista_existente(database):  # noqa: F811
    from scripts import migrate_20261020_02_vista_reporte_marcas as vista

    conn = database()
    c = conn.cursor()
    # Base vieja: sin la columna de correccion manual ni la tabla de correcciones, y con una vista propia.
    c.execute("CREATE TABLE asistencia_marcas (id BIGINT AUTO_INCREMENT PRIMARY KEY, fecha DATE)")
    c.execute("CREATE VIEW asistencia_marcas_reporte AS SELECT id, fecha, 'propia' AS origen FROM asistencia_marcas")
    for _ in range(2):  # se puede correr dos veces
        vista.migrate_cursor(c)
        conn.commit()
    c.execute("SELECT COUNT(*) FROM information_schema.COLUMNS WHERE TABLE_SCHEMA = DATABASE() "
              "AND TABLE_NAME = 'asistencia_marcas' AND COLUMN_NAME = 'corregida_manualmente'")
    assert c.fetchone()[0] == 1
    c.execute("SELECT COUNT(*) FROM asistencia_correcciones")
    assert c.fetchone()[0] == 0
    c.execute("SELECT VIEW_DEFINITION FROM information_schema.VIEWS WHERE TABLE_SCHEMA = DATABASE() "
              "AND TABLE_NAME = 'asistencia_marcas_reporte'")
    assert "propia" in c.fetchone()[0]  # la vista existente no se reemplaza
    c.close()
    conn.close()
