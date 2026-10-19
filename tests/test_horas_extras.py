"""Jornada laboral, etapa C: horas extras con autorizacion en dos pasos, topes, cierre del mes y liquidacion.

MySQL real (SH_TEST_DB_PORT). Usa los datos de la etapa B: el empleado 10 (horario L-V 08-12 y 13-17) reporta al 11."""
import datetime as dt
from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from services import horas_extras_service as he
from services import jornada_calculo_service as calc
from services import jornada_service as reglas
from tests.test_horas_jornada import _cargar_empleado_10, hj, jornada  # noqa: F401  (fixture)
from tests.test_jornada_laboral import QUIEN, _form, jl  # noqa: F401  (fixture)
from tests.test_seguridad import database  # noqa: F401  (fixture)

D = dt.date
RRHH = dict(empresa_id=1, user_id=99, empleado_id=None, aprobar=True, editar=True, exportar=True, **{"global": True})
JEFE = dict(empresa_id=1, user_id=98, empleado_id=11, aprobar=True, editar=False, exportar=False, **{"global": False})
SIN_EQUIPO = dict(JEFE, user_id=97, empleado_id=None)


@pytest.fixture
def he_db(hj):  # noqa: F811
    conn = hj()
    c = conn.cursor()
    for sql in ["ALTER TABLE empleados ADD COLUMN dni VARCHAR(20) NULL",
                "UPDATE empleados SET reporta_a_empleado_id = 11 WHERE id = 10",
                "UPDATE empleados SET reporta_a_empleado_id = 12 WHERE id = 11",  # si no tuviera jefe veria a toda la empresa
                "INSERT INTO usuarios VALUES (98, 'jefe'), (97, 'otro')"]:
        c.execute(sql)
    conn.commit()
    c.close()
    conn.close()
    return hj


def _fila(empleado_id, fecha):
    with calc.db.transaction(read_only=True) as c:
        return calc.db.one(c, "SELECT * FROM jl_horas_extra WHERE empleado_id = %s AND fecha = %s", (empleado_id, fecha))


# --- Clasificacion (sin base) -------------------------------------------------------------------------------------

def _t(dia, h1, h2, dia2=None):
    return (dt.datetime.combine(dia, dt.time.fromisoformat(h1)), dt.datetime.combine(dia2 or dia, dt.time.fromisoformat(h2)))


def test_clasificar_reparte_por_el_momento_en_que_ocurrio():
    sabado, domingo, martes, viernes = D(2026, 10, 10), D(2026, 10, 11), D(2026, 10, 6), D(2026, 10, 9)
    trece = dt.time(13, 0)
    assert calc.clasificar([_t(sabado, "09:00", "15:00")], 360, trece, {}) == (240, 120)  # sabado: 100 % desde las 13
    assert calc.clasificar([_t(martes, "08:00", "12:00"), _t(martes, "13:00", "19:00")], 120, trece, {}) == (120, 0)
    assert calc.clasificar([_t(domingo, "08:00", "10:00")], 120, trece, {}) == (0, 120)
    assert calc.clasificar([_t(martes, "08:00", "12:00")], 60, trece, {martes: "Feriado"}) == (0, 60)
    # Turno de viernes a sabado que termina antes de las 13: todo habil; los ultimos minutos son el excedente.
    assert calc.clasificar([_t(viernes, "22:00", "06:00", sabado)], 120, trece, {}) == (120, 0)
    assert calc.clasificar([_t(sabado, "12:00", "14:00")], 90, trece, {}) == (30, 60)
    assert calc.clasificar([_t(martes, "08:00", "09:00")], 75, trece, {}) == (75, 0)  # descanso computable: habil
    assert calc.clasificar([_t(sabado, "09:00", "15:00")], 360, None, {}) == (360, 0)  # sin hora de sabado: habil


# --- Dos pasos ----------------------------------------------------------------------------------------------------

def test_visto_bueno_del_jefe_y_decision_de_rrhh(he_db):
    _cargar_empleado_10()
    martes = D(2026, 10, 6)
    bandeja_jefe = he.bandeja(JEFE, "2026-10")
    assert {i["empleado"]["id"] for i in bandeja_jefe["items"]} == {10} and bandeja_jefe["alcance"] == "equipo"
    item = next(i for i in bandeja_jefe["items"] if i["dia"]["fecha"] == martes)
    assert (item["estado"], item["propuesta"], item["puede"]) == ("pendiente", (120, 0), ("jefe",))
    he.decidir(JEFE, 10, martes, "visto_bueno")
    assert (_fila(10, martes)["estado"], _fila(10, martes)["minutos_habil"]) == ("visto_jefe", 120)
    with pytest.raises(he.Error, match="más minutos que el excedente"):
        he.decidir(RRHH, 10, martes, "aprobar", "3:00", "0")
    he.decidir(RRHH, 10, martes, "aprobar", "1:00", "0:00")
    fila = _fila(10, martes)
    assert (fila["estado"], fila["minutos_habil"], fila["jefe_por"], fila["resuelto_por"]) == ("aprobada", 60, 98, 99)
    assert fila["recargo_habil"] == 50 and fila["recargo_descanso"] == 100  # de la LCT, guardados con la decision
    with pytest.raises(he.Error, match="No puede"):  # RR. HH. ya decidio: el jefe no lo cambia
        he.decidir(JEFE, 10, martes, "rechazo_jefe", motivo="No")
    detalle = he.detalle(RRHH, 10, martes)
    assert [e["accion"] for e in detalle["eventos"]] == ["visto_bueno", "aprobar"]
    assert detalle["eventos"][0]["usuario"] == "jefe"


def test_alcance_rechazos_y_reglas_de_decision(he_db):
    _cargar_empleado_10()
    jornada(11, "2026-10-06", ("08:00", "19:00"))  # el jefe tambien hizo excedente
    assert he.bandeja(SIN_EQUIPO, "2026-10")["items"] == []  # sin empleado vinculado: no ve a nadie
    with pytest.raises(he.Error, match="no está en su equipo"):
        he.decidir(JEFE, 11, D(2026, 10, 6), "visto_bueno")  # sus propias horas no
    with pytest.raises(he.Error, match="motivo del rechazo"):
        he.decidir(JEFE, 10, D(2026, 10, 12), "rechazo_jefe")
    he.decidir(JEFE, 10, D(2026, 10, 12), "rechazo_jefe", motivo="No estaba autorizado")
    assert he.estado_visible(_fila(10, D(2026, 10, 12))) == "rechazo_jefe"  # espera la decision de RR. HH.
    he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar")  # el feriado: todo con recargo de descanso
    assert (_fila(10, D(2026, 10, 12))["minutos_habil"], _fila(10, D(2026, 10, 12))["minutos_descanso"]) == (0, 240)
    he.decidir(RRHH, 10, D(2026, 10, 17), "aprobar")  # sabado de 9 a 11: antes de las 13, habil
    assert (_fila(10, D(2026, 10, 17))["minutos_habil"], _fila(10, D(2026, 10, 17))["minutos_descanso"]) == (120, 0)
    with pytest.raises(he.Error, match="no tiene excedente"):
        he.decidir(RRHH, 10, D(2026, 10, 8), "aprobar")  # dia a revisar
    hechos, errores = he.decidir_lote(RRHH, "2026-10", [(10, "2026-10-06"), (11, "2026-10-06"), (12, "2026-10-06")], "aprobar")
    assert hechos == 2 and len(errores) == 1  # el 12 no tiene excedente ese dia
    assert _fila(11, D(2026, 10, 6))["estado"] == "aprobada" and _fila(11, D(2026, 10, 6))["jefe_por"] is None


def test_superar_el_tope_exige_confirmar_y_motivo(he_db):
    reglas.guardar_regla(QUIEN, "empleado", 10, _form(vigente_desde="2026-10-01", tope_extras_mes="2"))
    _cargar_empleado_10()
    he.decidir(RRHH, 10, D(2026, 10, 6), "aprobar")  # 2 h: llega justo al tope
    item = next(i for i in he.bandeja(RRHH, "2026-10")["items"] if i["dia"]["fecha"] == D(2026, 10, 12))
    assert item["supera_tope"] and item["acumulado"][0] == 120
    with pytest.raises(he.Error, match="Supera el tope"):
        he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar")
    with pytest.raises(he.Error, match="Supera el tope"):
        he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar", confirmar_tope=True)  # sin motivo
    he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar", confirmar_tope=True, motivo="Inventario anual")
    assert _fila(10, D(2026, 10, 12))["supera_tope"] == 1
    he.decidir(JEFE, 10, D(2026, 10, 17), "visto_bueno")  # el jefe puede dar el visto bueno: decide RR. HH.
    assert _fila(10, D(2026, 10, 17))["supera_tope"] == 1


# --- Cierre y liquidacion -------------------------------------------------------------------------------------------

def test_cierre_liquidacion_y_reapertura(he_db, monkeypatch):
    jornada(10, "2026-10-05", ("08:00", "12:00"), ("13:00", "17:00"))
    jornada(10, "2026-10-06", ("08:00", "12:00"), ("13:00", "19:00"))
    with pytest.raises(he.Error, match="todavía no terminó"):
        he.cerrar(RRHH, "2026-10")
    monkeypatch.setattr(calc, "now_local", lambda: dt.datetime(2026, 11, 2, 9, 0))
    with pytest.raises(he.Error, match="1 días con excedente sin decisión"):
        he.cerrar(RRHH, "2026-10")
    he.decidir(JEFE, 10, D(2026, 10, 6), "visto_bueno")
    assert len(he.estado_cierre(RRHH, "2026-10")["pendientes"]) == 1  # el visto bueno no alcanza: decide RR. HH.
    he.decidir(RRHH, 10, D(2026, 10, 6), "aprobar")
    # Si despues se corrige el fichaje y el excedente baja, lo aprobado lo supera: no se cierra.
    with calc.db.transaction() as c:
        c.execute("UPDATE asistencia_marcas SET hora = '18:00' WHERE empleado_id = 10 AND fecha = '2026-10-06' AND hora = '19:00:00'")
    assert len(he.estado_cierre(RRHH, "2026-10")["excedidas"]) == 1
    with pytest.raises(he.Error, match="superan lo que hoy marcan"):
        he.cerrar(RRHH, "2026-10")
    he.decidir(RRHH, 10, D(2026, 10, 6), "aprobar")  # se decide de nuevo con la propuesta actual (1 h)
    with pytest.raises(he.Error, match="Solo RR. HH."):
        he.cerrar(JEFE, "2026-10")
    pid, liquidadas, _ = he.cerrar(RRHH, "2026-10")
    assert liquidadas == 1 and _fila(10, D(2026, 10, 6))["estado"] == "liquidada" and _fila(10, D(2026, 10, 6))["periodo_id"] == pid
    assert he.estado_cierre(RRHH, "2026-10")["cerrado"]
    with pytest.raises(he.Error, match="mes está cerrado"):
        he.decidir(RRHH, 10, D(2026, 10, 6), "rechazar", motivo="tarde")
    datos = he.liquidacion(RRHH, "2026-10")
    fila = next(f for f in datos["filas"] if f["id"] == 10)
    assert (fila["extras"]["habil"], fila["extras"]["recargos_habil"], int(fila["ausencias"])) == (60, {50}, 15)
    libro = load_workbook(BytesIO(he.exportar_liquidacion(datos)))
    assert libro.sheetnames == ["Liquidación", "Horas extras por día"]
    hoja = {r[1]: r for r in libro["Liquidación"].iter_rows(min_row=5, values_only=True)}
    assert hoja["Perez, Ana"][4] == 1.0 and hoja["Perez, Ana"][5] == "50 %"
    with pytest.raises(he.Error, match="motivo"):
        he.reabrir(RRHH, "2026-10", "")
    he.reabrir(RRHH, "2026-10", "Faltó una novedad")
    assert _fila(10, D(2026, 10, 6))["estado"] == "aprobada" and _fila(10, D(2026, 10, 6))["periodo_id"] is None
    assert [e["accion"] for e in he.detalle(RRHH, 10, D(2026, 10, 6))["eventos"]] == ["visto_bueno", "aprobar", "aprobar", "liquidar", "reabrir"]
    with pytest.raises(he.Error, match="mes cerrado"):
        he.liquidacion(RRHH, "2026-10")
    assert he.periodos(1)[0]["estado"] == "reabierto"


def test_cierre_bloqueado_por_dias_a_revisar(he_db, monkeypatch):
    _cargar_empleado_10()
    monkeypatch.setattr(calc, "now_local", lambda: dt.datetime(2026, 11, 2, 9, 0))
    estado = he.estado_cierre(RRHH, "2026-10")
    assert len(estado["revisar"]) == 3 and not estado["se_puede_cerrar"]


# --- Panel ------------------------------------------------------------------------------------------------------

@pytest.fixture
def client(he_db, monkeypatch):
    import web.auth.decorators as auth
    from web.horas import routes as horas_routes
    from web.horas_extras import routes

    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.secret_key = "test"
    app.config["TESTING"] = True
    app.register_blueprint(routes.web)
    app.register_blueprint(horas_routes.web)
    estado = {"acciones": {"ver", "aprobar", "editar", "exportar"}, "usuario": dict(id=99, empresa_id=1, activo=1, rol="rrhh")}

    def can_web(code, action="ver"):
        return code == "horas_extras" and action in estado["acciones"]

    app.jinja_env.globals.update(csp_nonce=lambda: "n", csrf_token=lambda: "t", can_web=can_web,
                                 export_toolbar={"enabled": False, "items": []})
    monkeypatch.setattr(auth, "can_access_module", lambda uid, module, action="ver": can_web(module, action))
    monkeypatch.setattr(routes, "can_access_module", lambda uid, module, action="ver": can_web(module, action))
    monkeypatch.setattr(routes, "_cached_web_user", lambda uid: estado["usuario"])
    monkeypatch.setattr(routes, "log_audit", lambda *a, **kw: None)
    test_client = app.test_client()
    with test_client.session_transaction() as sess:
        sess["user_id"] = 99
    test_client.estado = estado
    return test_client


def test_pantallas_y_permisos(client):
    _cargar_empleado_10()
    base = "/horas-extras"
    pagina = client.get(base + "/?mes=2026-10").get_data(as_text=True)
    assert "Perez, Ana" in pagina and "Aprobar la propuesta de los elegidos" in pagina and "Cierre del mes" in pagina
    dia = client.get(base + "/dia/10/2026-10-06").get_data(as_text=True)
    assert "Decisión de RR. HH." in dia and "Todavía no tiene visto bueno" in dia
    r = client.post(base + "/dia/10/2026-10-06", data={"accion": "aprobar", "minutos_habil": "1:30", "minutos_descanso": "0",
                                                       "volver": "https://otro.sitio/"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/horas-extras/dia/10/2026-10-06")  # nunca a otro sitio
    assert _fila(10, D(2026, 10, 6))["minutos_habil"] == 90
    r = client.post(base + "/lote", data={"mes": "2026-10", "accion": "aprobar", "dia": ["10|2026-10-12", "10|2026-10-17"]})
    assert r.status_code == 302 and _fila(10, D(2026, 10, 12))["estado"] == "aprobada"
    assert "Meses cerrados" not in client.get(base + "/cierre?mes=2026-10").get_data(as_text=True)
    assert "días con fichadas a revisar" in client.get(base + "/cierre?mes=2026-10").get_data(as_text=True)
    # El jefe: ve su equipo, da visto bueno, no ve el cierre.
    client.estado["usuario"] = dict(id=98, empresa_id=1, activo=1, rol="jefe", empleado_id=11)
    client.estado["acciones"] = {"ver", "aprobar"}
    pagina = client.get(base + "/?mes=2026-10&vista=todas").get_data(as_text=True)
    assert "Ves los días de tu equipo" in pagina and "Cierre del mes" not in pagina
    assert client.get(base + "/cierre?mes=2026-10").status_code == 403
    assert client.post(base + "/cierre/cerrar", data={"mes": "2026-10"}).status_code == 403
    client.estado["acciones"] = {"ver"}
    assert client.post(base + "/dia/10/2026-10-06", data={"accion": "visto_bueno"}).status_code == 403
    client.estado["acciones"] = set()
    assert client.get(base + "/").status_code == 403
