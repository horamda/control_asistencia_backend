import io
import pytest
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage
from services import seguridad_import_service as imp, seguridad_service as s
from tests.test_seguridad import database


def upload(names, status='OK', legajo=10):
    book=Workbook(); sheet=book.active; sheet.title='Respuestas'
    sheet.append(['Marca temporal','Que va a notificar','Apellido y nombre completo de la persona a informar',
                  'Legajo (informado)','Legajo (quien notifica)','Fecha del evento','Fecha del evento',
                  'Comportamiento Seguro','Comentarios, describir la tarea que se realizo'])
    for i,name in enumerate(names):
        sheet.append([f'2026-09-01T10:{i:02}:00','Comportamiento seguro',name,
                      'EXTERNO' if status=='EXTERNO' else legajo,11,None,'2026-09-01','EPP correcto','Usa EPP'])
    eq=book.create_sheet('Equivalencias')
    eq.append(['Nombre como fue escrito','Legajo','Empleado según lista','Estado','Nota'])
    for name in names: eq.append([name,legajo,'Visitante' if status=='EXTERNO' else 'Perez Ana',status,None])
    buffer=io.BytesIO();book.save(buffer);buffer.seek(0)
    return FileStorage(stream=buffer,filename='equivalencias.xlsx')


def test_approved_legajo_not_reporter_and_retry(database):
    file=upload(['Ana mal escrita']); raw=file.read();file.stream.seek(0)
    batch=imp.stage(1,99,file)
    assert imp.import_ready(1,batch,99)==(1,0)
    assert imp.stage(1,99,FileStorage(stream=io.BytesIO(raw),filename='otra_copia.xlsx'))==batch
    with s.db.transaction() as c:
        assert s.db.all_rows(c,'SELECT empleado_id FROM sh_involucrados')==[{'empleado_id':10}]
        assert s.db.one(c,'SELECT reportante_id FROM sh_eventos')['reportante_id'] is None


@pytest.mark.parametrize('status',['REVISAR','SIN IDENTIFICAR'])
def test_unapproved_equivalence_never_uses_filled_legajo(database,status):
    batch=imp.stage(1,99,upload(['Perez Ana'],status))
    assert imp.import_ready(1,batch,99)==(0,1)


def test_external_aliases_share_identity_and_export_fingerprint(database):
    from services import seguridad_excel_service as excel
    batch=imp.stage(1,99,upload(['Visitante','VISITANTE','Nombre alternativo'],'EXTERNO',None))
    assert imp.import_ready(1,batch,99)==(3,0)
    assert imp.import_ready(1,batch,99)==(0,0)
    with s.db.transaction() as c:
        assert len(s.external_people(c,1))==1
        assert s.db.one(c,'SELECT COUNT(DISTINCT externo_id) n FROM sh_evento_externos')['n']==1
    corrected=FileStorage(stream=io.BytesIO(excel.export(1,batch)),filename='corregido.xlsx')
    other=imp.stage(1,99,corrected)
    assert imp.import_ready(1,other,99)==(0,0)


def test_conflicting_legajo_is_pending(database):
    file=upload(['Persona'])
    from openpyxl import load_workbook
    book=load_workbook(file.stream);book['Respuestas']['D2']=11
    buffer=io.BytesIO();book.save(buffer);buffer.seek(0)
    batch=imp.stage(1,99,FileStorage(stream=buffer,filename='conflicto.xlsx'))
    assert imp.import_ready(1,batch,99)==(0,1)
