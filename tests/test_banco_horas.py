"""Jornada laboral, etapa D: banco de horas (acuerdos, libro que no se edita, vencimientos, francos y cierre del mes).

MySQL real (SH_TEST_DB_PORT). Usa los datos de las etapas B y C: el empleado 10 tiene horario L-V 08-12 y 13-17."""
import datetime as dt
from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import load_workbook

from services import banco_horas_service as bh
from services import horas_extras_service as he
from services import jornada_calculo_service as calc
from services import jornada_service as reglas
from tests.test_horas_extras import RRHH, _fila, he_db  # noqa: F401  (fixture)
from tests.test_horas_jornada import dias, hj, jornada, marcar  # noqa: F401  (fixture)
from tests.test_jornada_laboral import QUIEN, _form, jl  # noqa: F401  (fixture)
from tests.test_seguridad import database  # noqa: F401  (fixture)

D = dt.date


@pytest.fixture
def bh_db(he_db):  # noqa: F811
    """El empleado 10 tiene el banco habilitado por una condicion individual desde el 1/7/2026."""
    reglas.guardar_regla(QUIEN, "empleado", 10, _form(vigente_desde="2026-07-01", banco_horas_habilitado="1"))
    return he_db


def _mov(id_, fecha, minutos, vence=None, revierte=None):
    return dict(id=id_, fecha=fecha, minutos=minutos, vence=vence, revierte_id=revierte)


def _libro(empleado_id=10):
    with calc.db.transaction(read_only=True) as c:
        return calc.db.all_rows(c, "SELECT * FROM jl_banco_movimientos WHERE empleado_id = %s ORDER BY id", (empleado_id,))


# --- Lotes y saldo (sin base) ---------------------------------------------------------------------------------------

def test_los_debitos_consumen_primero_lo_que_vence_antes():
    movimientos = [_mov(1, D(2026, 1, 10), 120, D(2026, 12, 31)), _mov(2, D(2026, 2, 1), 60, D(2026, 6, 30)),
                   _mov(3, D(2026, 3, 1), 30), _mov(4, D(2026, 4, 1), -90)]
    assert [(l["id"], l["restante"]) for l in bh.lotes(movimientos)] == [(1, 90), (3, 30)]  # el 2 vencia antes
    estado = bh.saldo(movimientos, D(2026, 12, 15))
    assert (estado["total"], estado["disponible"], estado["por_vencer"]) == (120, 120, 90)
    assert bh.saldo(movimientos, D(2027, 1, 2))["vencido_sin_registrar"] == 90
    # Una reversion de un credito consume ese credito, no el que vence antes.
    assert [(l["id"], l["restante"]) for l in bh.lotes(movimientos + [_mov(5, D(2026, 4, 2), -30, revierte=3)])] == [(1, 90)]


def test_minutos_y_meses():
    assert (bh.minutos("2:30"), bh.minutos("2,5"), bh.minutos("90m")) == (150, 150, 90)
    with pytest.raises(bh.Error):
        bh.minutos("0")
    assert bh._mas_meses(D(2026, 1, 31), 1) == D(2026, 2, 28) and bh._mas_meses(D(2026, 11, 15), 3) == D(2027, 2, 15)


# --- Acuerdos ---------------------------------------------------------------------------------------------------

def test_acuerdo_requiere_la_regla_y_tiene_versiones(bh_db):
    with pytest.raises(bh.Error, match="no habilita el banco"):
        bh.guardar_acuerdo(RRHH, 11, _form(vigente_desde="2026-10-01"))
    primero = bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-08-01", saldo_maximo="40:00", vence_meses="6"))
    with calc.db.transaction(read_only=True) as c:
        a = bh.acuerdo_vigente(c, 10, D(2026, 9, 1))
    assert (str(a["factor_habil"]), str(a["factor_descanso"]), a["saldo_maximo_min"], a["vence_meses"]) == ("1.50", "2.00", 2400, 6)
    with pytest.raises(bh.Error, match="tiene que empezar después"):
        bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-07-15"))
    with pytest.raises(bh.Error, match="entre 0,5 y 3"):
        bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-10-01", factor_habil="5"))
    bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-10-01", factor_habil="1", factor_descanso="1", motivo="Una a una"))
    acuerdos = bh.detalle(RRHH, 10)["acuerdos"]
    assert [(x["vigente_desde"], x["vigente_hasta"]) for x in acuerdos] == [(D(2026, 10, 1), None), (D(2026, 8, 1), D(2026, 9, 30))]
    assert acuerdos[0]["reemplaza_id"] == primero
    bh.finalizar_acuerdo(RRHH, 10, "2026-12-31", "Fin del acuerdo")
    with pytest.raises(bh.Error, match="No tiene un acuerdo abierto"):
        bh.finalizar_acuerdo(RRHH, 10, "2027-01-31", "otra vez")


# --- Movimientos, francos y reversiones -----------------------------------------------------------------------------

def test_creditos_francos_y_reversiones(bh_db):
    with pytest.raises(bh.Error, match="acuerdo de banco de horas vigente"):
        bh.registrar(RRHH, 10, "credito", "2026-10-01", "8:00", "Saldo inicial")
    bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-09-01", saldo_maximo="10:00"))
    with pytest.raises(bh.Error, match="motivo"):
        bh.registrar(RRHH, 10, "credito", "2026-10-01", "8:00")
    with pytest.raises(bh.Error, match="saldo máximo"):
        bh.registrar(RRHH, 10, "credito", "2026-10-01", "12:00", "Saldo inicial")
    credito = bh.registrar(RRHH, 10, "credito", "2026-10-01", "8:00", "Saldo inicial")
    with pytest.raises(bh.Error, match="Saldo insuficiente"):
        bh.registrar(RRHH, 10, "debito", "2026-10-02", "9:00", "Ajuste")
    franco = bh.registrar(RRHH, 10, "franco", "2026-10-13", dia_completo=True)  # martes sin fichadas: 8 h programadas
    with pytest.raises(bh.Error, match="Ya tiene un franco"):
        bh.registrar(RRHH, 10, "franco", "2026-10-13", "1:00")
    with pytest.raises(bh.Error, match="Saldo insuficiente"):
        bh.registrar(RRHH, 10, "franco", "2026-10-21", "3:00")
    d, _ = dias(calc.calcular(1, D(2026, 10, 5), D(2026, 10, 23)), 10)
    assert (d[13]["estado"], d[13]["licencia"], d[13]["faltantes"]) == ("licencia", "Franco compensatorio", 0)
    with pytest.raises(bh.Error, match="ya se usaron"):
        bh.revertir(RRHH, credito, "Error de carga")  # esas horas ya se tomaron en el franco
    reversion = bh.revertir(RRHH, franco, "Vino a trabajar")
    with pytest.raises(bh.Error, match="ya fue revertido"):
        bh.revertir(RRHH, franco, "otra vez")
    with pytest.raises(bh.Error, match="no se revierte"):
        bh.revertir(RRHH, reversion, "otra vez")
    bh.registrar(RRHH, 10, "franco", "2026-10-21", "3:00")  # unas horas: baja lo programado
    d, _ = dias(calc.calcular(1, D(2026, 10, 5), D(2026, 10, 23)), 10)
    assert d[13]["estado"] == "ausente" and (d[21]["programados"], d[21]["faltantes"]) == (300, 300)
    assert bh.detalle(RRHH, 10)["saldo"]["total"] == 300
    assert [m["tipo"] for m in _libro()] == ["credito", "franco", "reversion", "franco"]  # nada se borra


def test_vencimiento_se_registra_en_el_libro(bh_db):
    bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-07-01", vence_meses="1"))
    bh.registrar(RRHH, 10, "credito", "2026-08-01", "4:00", "Inventario")
    estado = bh.detalle(RRHH, 10)["saldo"]
    assert (estado["total"], estado["disponible"], estado["vencido_sin_registrar"]) == (240, 0, 240)
    with pytest.raises(bh.Error, match="Saldo insuficiente"):
        bh.registrar(RRHH, 10, "debito", "2026-10-20", "1:00", "Ajuste")
    assert [m["tipo"] for m in _libro()] == ["credito"]  # el debito fallido no dejo nada
    bh.registrar(RRHH, 10, "credito", "2026-10-20", "1:00", "Guardia")  # el proximo movimiento registra lo vencido
    assert [(m["tipo"], m["minutos"]) for m in _libro()] == [("credito", 240), ("vencimiento", -240), ("credito", 60)]
    assert bh.detalle(RRHH, 10)["saldo"]["disponible"] == 60


# --- Horas extras al banco y cierre del mes -------------------------------------------------------------------------

def test_horas_extras_al_banco_faltantes_y_reapertura(bh_db, monkeypatch):
    jornada(10, "2026-10-05", ("08:00", "12:00"), ("13:00", "17:00"))
    jornada(10, "2026-10-06", ("08:00", "12:00"), ("13:00", "19:00"))
    jornada(10, "2026-10-07", ("08:05", "12:00"), ("13:00", "17:00"))  # 5 min faltantes
    jornada(10, "2026-10-12", ("08:00", "12:00"))  # feriado: 4 h de descanso
    with pytest.raises(bh.Error, match="solo se pueden pagar"):
        he.decidir(RRHH, 10, D(2026, 10, 6), "aprobar", destino="banco")
    bh.guardar_acuerdo(RRHH, 10, _form(vigente_desde="2026-10-01", saldo_maximo="10:00"))
    he.decidir(RRHH, 10, D(2026, 10, 6), "aprobar", destino="banco")  # 2 h habiles -> 3 h de banco
    with pytest.raises(bh.Error, match="saldo máximo"):
        he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar", "0", "4:00", destino="banco")  # 3 h + 4 h x 2 = 11 h > 10 h
    monkeypatch.setattr(calc, "now_local", lambda: dt.datetime(2026, 11, 2, 9, 0))
    he.decidir(RRHH, 10, D(2026, 10, 12), "aprobar")  # el feriado se paga
    pid, liquidadas, del_banco = he.cerrar(RRHH, "2026-10")
    assert liquidadas == 1 and {k: v for k, v in del_banco.items() if v} == dict(compensadas=1, acreditado=180, faltantes=5)
    assert _fila(10, D(2026, 10, 6))["estado"] == "compensada"
    assert [(m["tipo"], m["minutos"]) for m in _libro()] == [("extra", 180), ("faltantes", -5)]
    fila = next(f for f in he.liquidacion(RRHH, "2026-10")["filas"] if f["id"] == 10)
    assert (fila["extras"]["habil"], fila["extras"]["descanso"], fila["extras"]["banco"]) == (0, 240, 120)
    hoja = {r[1]: r for r in load_workbook(BytesIO(he.exportar_liquidacion(he.liquidacion(RRHH, "2026-10"))))["Liquidación"]
            .iter_rows(min_row=5, values_only=True)}
    assert hoja["Perez, Ana"][8] == 2.0  # horas extras al banco: no se pagan
    with pytest.raises(bh.Error, match="mes de esa fecha está cerrado"):
        bh.registrar(RRHH, 10, "franco", "2026-10-30", "1:00")
    bh.registrar(RRHH, 10, "franco", "2026-11-03", "2:00")
    with pytest.raises(bh.Error, match="quedaría en negativo"):
        he.reabrir(RRHH, "2026-10", "Corrección")
    assert he.estado_cierre(RRHH, "2026-10")["cerrado"]  # la reapertura fallida no cambio nada
    franco = next(m for m in _libro() if m["tipo"] == "franco")
    bh.revertir(RRHH, franco["id"], "No lo tomó")
    he.reabrir(RRHH, "2026-10", "Corrección")
    assert _fila(10, D(2026, 10, 6))["estado"] == "aprobada" and bh.detalle(RRHH, 10)["saldo"]["total"] == 0
    assert [e["accion"] for e in he.detalle(RRHH, 10, D(2026, 10, 6))["eventos"]] == ["aprobar", "compensar", "reabrir"]


# --- Panel ------------------------------------------------------------------------------------------------------

@pytest.fixture
def client(bh_db, monkeypatch):
    import web.auth.decorators as auth
    from web.banco_horas import routes
    from web.horas import routes as horas_routes  # filtro de horas del panel

    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.secret_key = "test"
    app.config["TESTING"] = True
    app.register_blueprint(routes.web)
    app.register_blueprint(horas_routes.web)
    permisos = {"acciones": {"ver", "editar"}}

    def can_web(code, action="ver"):
        return code == "banco_horas" and action in permisos["acciones"]

    app.jinja_env.globals.update(csp_nonce=lambda: "n", csrf_token=lambda: "t", can_web=can_web,
                                 export_toolbar={"enabled": False, "items": []})
    monkeypatch.setattr(auth, "can_access_module", lambda uid, module, action="ver": can_web(module, action))
    monkeypatch.setattr(routes, "_cached_web_user", lambda uid: dict(id=99, empresa_id=1, activo=1))
    monkeypatch.setattr(routes, "log_audit", lambda *a, **kw: None)
    test_client = app.test_client()
    with test_client.session_transaction() as sess:
        sess["user_id"] = 99
    test_client.permisos = permisos
    return test_client


def test_pantallas_y_permisos(client):
    base = "/banco-horas"
    assert "Todavía nadie tiene acuerdo" in client.get(base + "/").get_data(as_text=True)
    assert client.get(base + "/ir?empleado_id=10").headers["Location"].endswith("/banco-horas/empleado/10")
    pagina = client.get(base + "/empleado/10").get_data(as_text=True)
    assert "Cargar acuerdo" in pagina and "1,50" in pagina
    assert "no habilita el banco" in client.get(base + "/empleado/11").get_data(as_text=True)
    assert client.post(base + "/empleado/10/acuerdo", data={"vigente_desde": "2026-10-01"}).status_code == 302
    r = client.post(base + "/empleado/10/movimiento", data={"tipo": "credito", "fecha": "2026-10-01", "horas": "6:00", "motivo": "Inicial"})
    assert r.status_code == 302 and bh.detalle(RRHH, 10)["saldo"]["total"] == 360
    assert "Perez, Ana" in client.get(base + "/").get_data(as_text=True)
    movimiento = _libro()[0]["id"]
    assert client.post(f"{base}/movimiento/{movimiento}/revertir", data={"motivo": "Error"}).status_code == 302
    assert bh.detalle(RRHH, 10)["saldo"]["total"] == 0 and "revertido" in client.get(base + "/empleado/10").get_data(as_text=True)
    client.permisos["acciones"] = {"ver"}
    assert client.post(base + "/empleado/10/movimiento", data={"tipo": "debito"}).status_code == 403
    assert "Cargar un movimiento" not in client.get(base + "/empleado/10").get_data(as_text=True)
    assert "Cargar un movimiento" not in client.get(base + "/empleado/11").get_data(as_text=True)
    client.permisos["acciones"] = set()
    assert client.get(base + "/").status_code == 403
