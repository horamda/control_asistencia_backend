import io
import json
from datetime import date, time
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

from services import seguridad_excel_service as excel
from services import seguridad_excel_schema as schema
from services import seguridad_import_service as imp
from services import seguridad_service as s
from tests.test_seguridad import database, client


def upload(rows, name='plantilla.xlsx'):
    book=Workbook(); sheet=book.active; sheet.title='Carga'
    sheet.append(schema.HEADERS)
    for row in rows:
        sheet.append([row.get(h) for h in schema.HEADERS])
    buffer=io.BytesIO(); book.save(buffer); buffer.seek(0)
    return FileStorage(stream=buffer,filename=name)


def complete(**changes):
    row={'ID registro':str(uuid4()),'Legajos':'0010','Tipo':'seguro',
         'Fecha':date(2026,9,1),'Categoría':'EPP correcto','Descripción':'Usa protección','Tipo persona':'empleado'}
    row.update(changes)
    return row


def test_template_roundtrip_and_multiple_people(database):
    raw=excel.export(1)
    book=load_workbook(io.BytesIO(raw))
    assert book['Carga']['C2'].number_format=='@'
    assert len(book['Carga'].data_validations.dataValidation)==6
    assert book['Carga']['Q2'].data_type=='f'
    assert '0010' in [r[0] for r in book['Empleados'].iter_rows(min_row=2,values_only=True)]
    with pytest.raises(s.Error,match='respuestas'):
        imp.read_file(FileStorage(stream=io.BytesIO(raw),filename='vacio.xlsx'))
    row=complete(Legajos='0010;0011')
    bid=imp.stage(1,99,upload([row]))
    assert imp.import_ready(1,bid,99)==(1,0)
    event=s.history(1)['items'][0]
    assert event['estado']=='pendiente'
    assert len(s.detail(1,event['id'])['involucrados'])==2
    # Changing workbook bytes/layout must not duplicate the same ID.
    row['Persona de referencia']='Nombre de referencia'
    second=imp.stage(1,99,upload([row]))
    assert imp.import_ready(1,second,99)==(0,0)
    assert s.history(1)['total']==1
    row['Descripción']='Contenido diferente'
    with pytest.raises(imp.ValidationError,match='errores'):
        imp.stage(1,99,upload([row]))
    assert s.history(1)['total']==1


@pytest.mark.parametrize('changes,field',[
    ({'Legajos':'10'},'Legajos'), # Preserve leading zeroes; no numeric guess.
    ({'Legajos':'0020'},'Legajos'), # Other company.
    ({'Legajos':'0010;0010'},'Legajos'),
    ({'Fecha':'2030-01-01'},'Fecha'),
    ({'Categoría':'Falta EPP'},'Categoría'),
    ({'Descripción':'=HYPERLINK("https://example.com")'},'Descripción'),
    ({'ID registro':'no-id'},'ID registro'),
    ({'Tipo':'inseguro','Categoría':'Falta EPP','Advertido':''},'Advertido'),
    ({'Tipo':'accidente','Categoría':'','Hora':'','Lugar':''},'Hora'),
])
def test_rejects_entire_template_before_saving(database,changes,field):
    with pytest.raises(imp.ValidationError) as exc:
        imp.stage(1,99,upload([complete(),complete(**changes)]))
    assert exc.value.issues[0]['fila']==3
    assert any(field in error for error in exc.value.issues[0]['errores'])
    with s.db.transaction() as c:
        assert s.db.one(c,'SELECT COUNT(*) n FROM sh_importaciones')['n']==0
        assert s.db.one(c,'SELECT COUNT(*) n FROM sh_eventos')['n']==0


def test_duplicate_ids_and_excel_times(database):
    row=complete()
    with pytest.raises(imp.ValidationError):
        imp.stage(1,99,upload([row,row]))
    row=complete(Tipo='incidente',Categoría='',Hora=time(11,20),Lugar='Depósito')
    bid=imp.stage(1,99,upload([row]))
    assert imp.import_ready(1,bid,99)==(1,0)


def test_empty_historical_rows_are_not_silently_skipped():
    file=upload([{'ID registro':str(uuid4()),'Fila origen':77}])
    _,rows=imp.read_file(file)
    assert len(rows)==1
    _,errors=imp.suggest(rows[0]['original'],[],[])
    assert any('Tipo persona:' in e for e in errors)
    assert any('Descripción:' in e for e in errors)


def test_correcting_legacy_links_original_and_preserves_bytes(database,monkeypatch):
    headers=['Marca temporal','Que va a notificar','Nombre Completo de quien comete la infracción',
             'Fecha que se observa la condición insegura-comportamiento inseguro','Infracción que cometió',
             'Comentarios','Se advirtió a la persona que estaba cometiendo el acto?']
    original={'encabezados':headers,'celdas':['2026-09-01','Comportamiento Inseguro','Nombre ambiguo',
              '2026-09-01','Falta EPP','Sin protección','Sí']}
    with monkeypatch.context() as patch:
        patch.setattr(imp,'read_file',lambda f:(b'original intacto',[dict(hoja='Form',fila=3,original=original)]))
        bid=imp.stage(1,99,FileStorage(filename='original.xlsx'))
    exported=excel.export(1,bid)
    book=load_workbook(io.BytesIO(exported))
    sheet=book['Carga']
    assert sheet['C2'].value is None
    assert sheet['A2'].value==imp.source_id(1,original)
    source=sheet['B2'].value
    sheet['C2']='0010'
    sheet['R2']='empleado'
    buf=io.BytesIO(); book.save(buf); buf.seek(0)
    corrected=imp.stage(1,99,FileStorage(stream=buf,filename='corregido.xlsx'))
    assert imp.import_ready(1,corrected,99)==(1,0)
    assert imp.batch(1,bid)['totales']['resueltas']==1
    assert imp.import_ready(1,bid,99)==(0,0)
    with s.db.transaction() as c:
        assert s.db.one(c,'SELECT archivo FROM sh_importaciones WHERE id=%s',(bid,))['archivo']==b'original intacto'
        assert json.loads(s.db.one(c,'SELECT original FROM sh_import_filas WHERE id=%s',(source,))['original'])==original
    # Tampering with the source must fail, even with otherwise valid data.
    with pytest.raises(imp.ValidationError):
        imp.stage(1,99,upload([complete(**{'Fila origen':source})]))
    with pytest.raises(s.Error):
        excel.export(2,bid)


def test_download_permissions_and_error_screen(client,monkeypatch):
    import web.auth.decorators as auth
    response=client.get('/seguridad-higiene/importar/plantilla')
    assert response.status_code==200
    assert response.headers['Cache-Control']=='private, no-store'
    bad=upload([complete(Legajos='invalid')])
    response=client.post('/seguridad-higiene/importar',data={'archivo':(bad.stream,bad.filename)})
    assert response.status_code==400
    assert 'No se guardó la plantilla' in response.get_data(as_text=True)
    assert 'Legajos:' in response.get_data(as_text=True)
    monkeypatch.setattr(auth,'can_access_module',lambda uid,module,action='ver':action=='ver')
    assert client.get('/seguridad-higiene/importar/plantilla').status_code==403
    assert client.get('/seguridad-higiene/importar/1/corregir').status_code==403


def test_export_neutralizes_formulas(database):
    with s.db.transaction() as c:
        people=s.employees(c,1,False)
        cats=s.db.all_rows(c,'SELECT * FROM sh_catalogos WHERE empresa_id=1')
    people[0]['apellido']='=HYPERLINK("https://example.com")'
    book=load_workbook(io.BytesIO(excel.build(1,people,cats)))
    assert book['Empleados']['B2'].data_type=='s'
