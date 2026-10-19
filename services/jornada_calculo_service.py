"""Motor de calculo de horas (jornada laboral, etapa B). Solo LEE fichajes: nunca los modifica.

Por empleado y dia calcula, en minutos:
- programados: los bloques del horario asignado (con excepciones y feriados); sin horario, las horas diarias pactadas
  en los dias laborales de su regla.
- fichados: suma de pares ingreso -> egreso. Un ingreso se empareja con el proximo egreso hasta MAX_TURNO despues,
  aunque sea al dia siguiente (turno que cruza la medianoche): el turno cuenta para el dia del ingreso.
- trabajados: fichados + el descanso computable de su regla (hasta lo que realmente descanso entre pares).
- excedente (trabajados - programados), faltantes (programados - trabajados, solo dias laborables cerrados y sin
  licencia), ordinarios (trabajados - excedente), nocturnos (dentro de la franja nocturna de su regla) y descansos.

No clasifica el excedente como hora extra legal (eso es la etapa C) ni lo autoriza. Las tolerancias del horario no se
usan: nunca cambian los minutos trabajados. Un dia con fichajes incompletos o cerrado por el egreso automatico queda
"a revisar" sin excedente ni faltante.

El calculo se hace en vivo; "registrar" guarda una foto del periodo (jl_calculos + jl_resultados_dia) con las
versiones de regla usadas y una huella de las marcas, conservando los calculos anteriores."""

import datetime as dt
import hashlib
import json
from collections import defaultdict

from repositories import carga_repository as db
from services import feriado_service
from services import jornada_service as reglas
from services.carga_service import CargaError as Error, now_local

MAX_TURNO = dt.timedelta(hours=16)
MAX_DIAS = 62
JORNADA_LARGA_MIN = 12 * 60
DESCANSO_ENTRE_JORNADAS_MIN = 12 * 60  # LCT art. 197: 12 h entre una jornada y la siguiente
ANULA = {"FRANCO", "VACACIONES", "FERIADO", "LICENCIA"}

ESTADOS = {
    "normal": "Normal",
    "revisar": "A revisar",
    "ausente": "Ausente",
    "licencia": "Licencia o justificada",
    "descanso": "Día no laborable",
    "pendiente": "Día en curso o futuro",
    "fuera": "Fuera de la relación laboral",
}
DIAS_CORTOS = ("Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom")
TIPOS_DIA = {"laborable": "Laborable", "no_laborable": "No laborable", "feriado": "Feriado", "licencia": "Licencia",
             "fuera": "Fuera de la relación laboral"}


# --- Utilidades ---------------------------------------------------------------------------------------------

def _hoy():
    return now_local().date()


def _minutos_hora(valor):
    if valor is None:
        return None
    if isinstance(valor, dt.timedelta):
        return int(valor.total_seconds() // 60)
    if hasattr(valor, "hour"):
        return valor.hour * 60 + valor.minute
    h, m = str(valor).split(":")[:2]
    return int(h) * 60 + int(m)


def _fecha(valor):
    return valor if isinstance(valor, dt.date) else dt.date.fromisoformat(str(valor)[:10])


def _momento(fecha, hora):
    minutos = _minutos_hora(hora)
    return dt.datetime.combine(_fecha(fecha), dt.time()) + dt.timedelta(minutes=minutos)


def _hhmm(minutos):
    return f"{minutos // 60:02d}:{minutos % 60:02d}"


def horas(minutos):
    """Minutos -> "8:30" para mostrar (los calculos siempre en minutos)."""
    minutos = int(minutos or 0)
    signo = "-" if minutos < 0 else ""
    minutos = abs(minutos)
    return f"{signo}{minutos // 60}:{minutos % 60:02d}"


def _solape(a_ini, a_fin, b_ini, b_fin):
    return max(0, int((min(a_fin, b_fin) - max(a_ini, b_ini)).total_seconds() // 60))


# --- Pares de fichadas -----------------------------------------------------------------------------------------

def emparejar(marcas):
    """marcas: [{id, fecha, hora, accion, metodo}] de un empleado. Devuelve ({fecha: [pares]}, {fecha: [sueltas]}).
    Un par es (inicio, fin, automatico) y cuenta para la fecha de su ingreso."""
    ordenadas = sorted(marcas, key=lambda m: (_momento(m["fecha"], m["hora"]), int(m.get("id") or 0)))
    pares, sueltas = defaultdict(list), defaultdict(list)
    abierto = None
    for marca in ordenadas:
        momento = _momento(marca["fecha"], marca["hora"])
        accion = str(marca.get("accion") or "").lower()
        if accion == "ingreso":
            if abierto:
                sueltas[abierto["fecha"]].append(abierto)
            abierto = dict(marca, momento=momento, fecha=_fecha(marca["fecha"]))
        elif accion == "egreso":
            if abierto and momento - abierto["momento"] <= MAX_TURNO and momento > abierto["momento"]:
                automatico = str(marca.get("metodo") or "").lower().startswith("auto")
                pares[abierto["fecha"]].append((abierto["momento"], momento, automatico))
                abierto = None
            else:
                if abierto:
                    sueltas[abierto["fecha"]].append(abierto)
                    abierto = None
                sueltas[_fecha(marca["fecha"])].append(dict(marca, momento=momento))
        else:
            sueltas[_fecha(marca["fecha"])].append(dict(marca, momento=momento))
    if abierto:
        sueltas[abierto["fecha"]].append(abierto)
    return pares, sueltas


# --- Horarios por lote ------------------------------------------------------------------------------------------

class Horarios:
    """Horario esperado de muchos empleados y dias con pocas consultas (misma logica que get_horario_esperado:
    una excepcion que anula manda; un CAMBIO_HORARIO usa sus bloques; si no, los del dia del horario asignado)."""

    def __init__(self, c, empleado_ids, desde, hasta):
        ids = sorted(empleado_ids)
        marcas = ",".join(["%s"] * len(ids)) or "NULL"
        self.asignaciones = defaultdict(list)
        for a in db.all_rows(c, f"""SELECT eh.empleado_id, eh.horario_id, eh.fecha_desde, eh.fecha_hasta, h.nombre,
                                           h.tolerancia_min
                                    FROM empleado_horarios eh JOIN horarios h ON h.id = eh.horario_id
                                    WHERE eh.empleado_id IN ({marcas}) AND eh.fecha_desde <= %s
                                      AND (eh.fecha_hasta IS NULL OR eh.fecha_hasta >= %s)
                                    ORDER BY eh.fecha_desde DESC""", (*ids, hasta, desde)) if ids else []:
            self.asignaciones[int(a["empleado_id"])].append(a)
        horario_ids = sorted({int(a["horario_id"]) for lista in self.asignaciones.values() for a in lista})
        self.dias = defaultdict(dict)  # horario -> {dia_semana: [(entrada, salida)]}
        if horario_ids:
            mh = ",".join(["%s"] * len(horario_ids))
            for fila in db.all_rows(c, f"""SELECT d.horario_id, d.dia_semana, b.hora_entrada, b.hora_salida
                                           FROM horario_dias d LEFT JOIN horario_dia_bloques b ON b.horario_dia_id = d.id
                                           WHERE d.horario_id IN ({mh}) ORDER BY d.horario_id, d.dia_semana, b.orden""",
                                    tuple(horario_ids)):
                bloques = self.dias[int(fila["horario_id"])].setdefault(int(fila["dia_semana"]), [])
                if fila["hora_entrada"] is not None and fila["hora_salida"] is not None:
                    bloques.append((_minutos_hora(fila["hora_entrada"]), _minutos_hora(fila["hora_salida"])))
        self.excepciones = {}
        if ids:
            for e in db.all_rows(c, f"""SELECT x.id, x.empleado_id, x.fecha, x.tipo, x.anula_horario, b.hora_entrada, b.hora_salida
                                        FROM empleado_excepciones x LEFT JOIN excepcion_bloques b ON b.excepcion_id = x.id
                                        WHERE x.empleado_id IN ({marcas}) AND x.fecha BETWEEN %s AND %s
                                        ORDER BY x.id, b.orden""", (*ids, desde, hasta)):
                clave = (int(e["empleado_id"]), _fecha(e["fecha"]))
                actual = self.excepciones.setdefault(clave, dict(tipo=e["tipo"], anula=bool(e["anula_horario"]), bloques=[]))
                if e["hora_entrada"] is not None and e["hora_salida"] is not None:
                    actual["bloques"].append((_minutos_hora(e["hora_entrada"]), _minutos_hora(e["hora_salida"])))

    def esperado(self, empleado_id, fecha):
        """dict(horario, bloques[(min, min)], anula, tipo_excepcion) o None si no tiene horario ni excepcion."""
        excepcion = self.excepciones.get((int(empleado_id), fecha))
        asignado = next((a for a in self.asignaciones.get(int(empleado_id), [])
                         if _fecha(a["fecha_desde"]) <= fecha and (a["fecha_hasta"] is None or _fecha(a["fecha_hasta"]) >= fecha)), None)
        if excepcion:
            tipo = str(excepcion["tipo"] or "").upper()
            if excepcion["anula"] or tipo in ANULA:
                return dict(horario=asignado and asignado["nombre"], bloques=excepcion["bloques"], anula=True, tipo_excepcion=tipo)
            if tipo == "CAMBIO_HORARIO":
                return dict(horario=asignado and asignado["nombre"], bloques=excepcion["bloques"], anula=False, tipo_excepcion=tipo)
        if not asignado:
            return None
        dias = self.dias.get(int(asignado["horario_id"]), {})
        bloques = dias.get(fecha.isoweekday(), []) if dias else []
        return dict(horario=asignado["nombre"], bloques=bloques, anula=False, tipo_excepcion=None)


# --- Calculo ----------------------------------------------------------------------------------------------------

def _nocturnos(pares, regla):
    desde, hasta = regla.get("nocturna_desde"), regla.get("nocturna_hasta")
    if desde is None or hasta is None:
        return 0
    d, h = _minutos_hora(desde), _minutos_hora(hasta)
    total = 0
    for inicio, fin, _ in pares:
        dia = dt.datetime.combine(inicio.date(), dt.time())
        for base in (dia - dt.timedelta(days=1), dia, dia + dt.timedelta(days=1)):
            ini = base + dt.timedelta(minutes=d)
            fin_franja = base + dt.timedelta(minutes=h if h > d else h + 24 * 60)
            total += _solape(inicio, fin, ini, fin_franja)
    return total


def _dia(fecha, pares, sueltas, esperado, regla, feriado, licencia, hoy, vinculado=True, franco=0):
    """Resultado de un dia (minutos y estado) con su detalle. vinculado: la fecha esta dentro de la relacion laboral.
    franco: minutos de franco compensatorio del banco de horas (el dia completo es licencia; si no, baja lo programado)."""
    pactadas = regla.get("horas_diarias")
    dias_laborales = [int(x) for x in str(regla.get("dias_laborales") or "").split(",") if x.strip().isdigit()]
    trabaja_feriados = bool(regla.get("trabaja_feriados"))
    if esperado:
        bloques = [] if esperado["anula"] else esperado["bloques"]
        programados = sum(max(0, f - i) for i, f in bloques)
    else:
        bloques = []
        programados = int(float(pactadas) * 60) if pactadas and fecha.isoweekday() in dias_laborales else 0
    tipo_dia = "laborable" if programados else "no_laborable"
    if feriado and not trabaja_feriados:
        tipo_dia, programados = "feriado", 0
    if licencia or (esperado and esperado["anula"]):
        tipo_dia, programados = "licencia", 0
    if franco and tipo_dia == "laborable":
        if franco >= programados:
            tipo_dia, programados, licencia = "licencia", 0, licencia or "Franco compensatorio"
        else:
            programados -= franco
    if not vinculado:
        tipo_dia, programados = "fuera", 0

    fichados = sum(int((f - i).total_seconds() // 60) for i, f, _ in pares)
    ordenados = sorted(pares)
    descanso = sum(max(0, int((ordenados[k + 1][0] - ordenados[k][1]).total_seconds() // 60)) for k in range(len(ordenados) - 1))
    computable = min(descanso, int(regla.get("descanso_computable_min") or 0))
    trabajados = fichados + computable
    automatico = any(a for _, _, a in pares)
    incompleto = bool(sueltas) or automatico
    excedente = faltantes = 0
    if fecha >= hoy:
        estado = "pendiente"
    elif incompleto or (pares and tipo_dia in ("licencia", "fuera")):
        estado = "revisar"
    elif tipo_dia == "fuera":
        estado = "fuera"
    elif tipo_dia == "licencia":
        estado = "licencia"
    elif not pares and programados:
        estado = "ausente"
    elif not pares:
        estado = "descanso"
    else:
        estado = "normal"
    if estado == "normal":
        excedente = max(0, trabajados - programados)
        faltantes = max(0, programados - trabajados) if tipo_dia == "laborable" else 0
    elif estado == "ausente":
        faltantes = programados
    ordinarios = trabajados - excedente
    motivo = None
    if automatico:
        motivo = "Cerrado por el egreso automático: confirmar la hora real de salida."
    elif sueltas:
        motivo = "Fichadas sin pareja (falta un ingreso o un egreso)."
    elif pares and tipo_dia == "licencia":
        motivo = f"Fichó en un día de {(licencia or (esperado or {}).get('tipo_excepcion') or 'licencia').lower()}."
    elif pares and tipo_dia == "fuera":
        motivo = "Fichó fuera de la relación laboral (antes del ingreso o después de la baja)."
    return dict(
        fecha=fecha, estado=estado, tipo_dia=tipo_dia, programados=programados, fichados=fichados, trabajados=trabajados,
        ordinarios=ordinarios, excedente=excedente, faltantes=faltantes, nocturnos=_nocturnos(pares, regla),
        descanso=descanso, cruza_medianoche=any(f.date() > i.date() for i, f, _ in pares), motivo=motivo,
        feriado=feriado, licencia=licencia, horario=(esperado or {}).get("horario"), tipo_excepcion=(esperado or {}).get("tipo_excepcion"),
        bloques=[f"{_hhmm(i)}–{_hhmm(f)}" for i, f in bloques],
        pares=[f"{i:%H:%M}–{f:%H:%M}" + (" (+1)" if f.date() > i.date() else "") + (" auto" if a else "")
               for i, f, a in ordenados],
        sueltas=[f"{s['accion']} {s['momento']:%H:%M}" for s in sueltas],
        tramos=[(i, f) for i, f, _ in ordenados],
    )


def clasificar(tramos, excedente, sabado_desde, feriados):
    """Propuesta de reparto del excedente (nunca una decision): (minutos con recargo habil, con recargo de descanso).

    Se toman como excedente los ultimos minutos trabajados del dia y cada tramo se clasifica por el momento en que
    ocurrio (LCT art. 201): descanso si es domingo, feriado o sabado desde la hora de la regla; habil el resto."""
    if excedente <= 0:
        return (0, 0)
    corte_sabado = _minutos_hora(sabado_desde)
    segmentos, resta = [], excedente
    for inicio, fin in sorted(tramos, reverse=True):
        if resta <= 0:
            break
        toma = min(resta, int((fin - inicio).total_seconds() // 60))
        segmentos.append((fin - dt.timedelta(minutes=toma), fin))
        resta -= toma
    habil, descanso = resta, 0  # lo que no alcanza a cubrir lo fichado (descanso computable) va como habil
    for a, b in segmentos:
        cortes = {a, b}
        dia = a.date()
        while dia <= b.date():
            base = dt.datetime.combine(dia, dt.time())
            cortes.add(base)
            if corte_sabado is not None and dia.isoweekday() == 6:
                cortes.add(base + dt.timedelta(minutes=corte_sabado))
            dia += dt.timedelta(days=1)
        puntos = sorted(x for x in cortes if a <= x <= b)
        for x, y in zip(puntos, puntos[1:]):
            minutos = int((y - x).total_seconds() // 60)
            d = x.date()
            es_descanso = (d in feriados or d.isoweekday() == 7 or
                           (d.isoweekday() == 6 and corte_sabado is not None and x.hour * 60 + x.minute >= corte_sabado))
            if es_descanso:
                descanso += minutos
            else:
                habil += minutos
    return (habil, descanso)


def _empleados(c, empresa_id, filtros, desde, hasta):
    """Activos y los dados de baja dentro del periodo (sus dias fuera de la relacion laboral no cuentan)."""
    donde, args = ["e.empresa_id = %s", "(e.activo = 1 OR e.fecha_baja >= %s)"], [empresa_id, desde]
    for clave, columna in (("sucursal_id", "e.sucursal_id"), ("sector_id", "e.sector_id"), ("empleado_id", "e.id"),
                           ("jefe_id", "e.reporta_a_empleado_id")):
        valor = str(filtros.get(clave) or "")
        if valor.isdigit():
            donde.append(f"{columna} = %s")
            args.append(int(valor))
    ids = filtros.get("ids")
    if ids is not None:  # alcance de un jefe: solo su cadena
        ids = sorted({int(x) for x in ids})
        if not ids:
            return []
        donde.append(f"e.id IN ({','.join(['%s'] * len(ids))})")
        args.extend(ids)
    convenio = str(filtros.get("convenio_id") or "")
    if convenio.isdigit():
        donde.append("""e.id IN (SELECT a.empleado_id FROM jl_asignaciones a WHERE a.convenio_id = %s
                         AND a.vigente_desde <= %s AND (a.vigente_hasta IS NULL OR a.vigente_hasta >= %s))""")
        args.extend([int(convenio), hasta, desde])
    return db.all_rows(c, f"""SELECT e.id, e.apellido, e.nombre, e.legajo, e.sucursal_id, e.sector_id, e.activo,
                                     e.fecha_ingreso, e.fecha_baja, su.nombre AS sucursal_nombre, se.nombre AS sector_nombre
                              FROM empleados e LEFT JOIN sucursales su ON su.id = e.sucursal_id
                              LEFT JOIN sectores se ON se.id = e.sector_id
                              WHERE {' AND '.join(donde)} ORDER BY e.apellido, e.nombre""", tuple(args))


def _vinculado(empleado, fecha):
    ingreso, baja = empleado.get("fecha_ingreso"), empleado.get("fecha_baja")
    return (not ingreso or _fecha(ingreso) <= fecha) and (not baja or _fecha(baja) >= fecha)


def _rango(desde, hasta):
    desde, hasta = _fecha(desde), _fecha(hasta)
    if desde > hasta:
        raise Error("'Desde' no puede ser posterior a 'Hasta'.")
    if (hasta - desde).days + 1 > MAX_DIAS:
        raise Error(f"El período máximo es de {MAX_DIAS} días.")
    return desde, hasta


def calcular(empresa_id, desde, hasta, filtros=None):
    """Calculo en vivo del periodo. Devuelve {empleados: [{empleado, dias: [...], totales, semanas, alertas}], totales}."""
    filtros = filtros or {}
    desde, hasta = _rango(desde, hasta)
    hoy = _hoy()
    with db.transaction(read_only=True) as c:
        empleados = _empleados(c, empresa_id, filtros, desde, hasta)
        ids = [int(e["id"]) for e in empleados]
        marcas = ",".join(["%s"] * len(ids)) or "NULL"
        filas = db.all_rows(c, f"""SELECT id, empleado_id, fecha, hora, accion, metodo FROM asistencia_marcas
                                   WHERE empresa_id = %s AND empleado_id IN ({marcas}) AND fecha BETWEEN %s AND %s""",
                            (empresa_id, *ids, desde - dt.timedelta(days=1), hasta + dt.timedelta(days=1))) if ids else []
        horarios = Horarios(c, ids, desde, hasta)
        feriados = feriado_service.mapa(c, empresa_id)
        licencias = _licencias(c, ids, desde, hasta)
        francos = _francos_banco(c, ids, desde, hasta)
    resolutor = reglas.Resolutor(empresa_id, empleados, desde, hasta)
    por_empleado = defaultdict(list)
    for f in filas:
        por_empleado[int(f["empleado_id"])].append(f)
    resultado = []
    for e in empleados:
        eid = int(e["id"])
        pares, sueltas = emparejar(por_empleado.get(eid, []))
        feriados_suc = feriado_service.de_sucursal(feriados, e.get("sucursal_id") or 0)
        dias = []
        fecha = desde
        while fecha <= hasta:
            regla, usadas, _ = resolutor.efectivas(eid, fecha)
            dia = _dia(fecha, pares.get(fecha, []), sueltas.get(fecha, []), horarios.esperado(eid, fecha), regla,
                       feriados_suc.get(fecha), licencias.get((eid, fecha)), hoy, _vinculado(e, fecha),
                       francos.get((eid, fecha), 0))
            dia["reglas"] = usadas
            dia["huella"] = _huella(por_empleado.get(eid, []), fecha)
            dia["pactadas_diarias"] = int(float(regla["horas_diarias"]) * 60) if regla.get("horas_diarias") else None
            dia["recargos"] = (regla.get("recargo_habil_pct"), regla.get("recargo_descanso_pct"))
            dia["topes"] = tuple(int(regla[k]) * 60 if regla.get(k) not in (None, "") else None
                                 for k in ("tope_extras_mes", "tope_extras_anio"))
            dia["propuesta"] = clasificar(dia["tramos"], dia["excedente"], regla.get("sabado_descanso_desde"), feriados_suc)
            dias.append(dia)
            fecha += dt.timedelta(days=1)
        resultado.append(dict(empleado=e, dias=dias, totales=_totales(dias), semanas=_semanas(dias, resolutor, eid),
                              alertas=_alertas(dias, pares)))
    return dict(desde=desde, hasta=hasta, empleados=resultado, totales=_totales([d for r in resultado for d in r["dias"]]),
                calculado_at=now_local().replace(microsecond=0))


def _licencias(c, ids, desde, hasta):
    """{(empleado, fecha): motivo}: justificaciones aprobadas, vacaciones aprobadas y francos."""
    salida = {}
    if not ids:
        return salida
    marcas = ",".join(["%s"] * len(ids))

    def cubrir(eid, d1, d2, motivo):
        d1, d2 = max(_fecha(d1), desde), min(_fecha(d2), hasta)
        while d1 <= d2:
            salida.setdefault((int(eid), d1), motivo)
            d1 += dt.timedelta(days=1)

    for j in db.all_rows(c, f"""SELECT empleado_id, fecha, fecha_desde, fecha_hasta FROM justificaciones
                                WHERE estado = 'aprobada' AND empleado_id IN ({marcas})
                                  AND COALESCE(fecha_hasta, fecha) >= %s AND COALESCE(fecha_desde, fecha) <= %s""",
                         (*ids, desde, hasta)):
        cubrir(j["empleado_id"], j["fecha_desde"] or j["fecha"], j["fecha_hasta"] or j["fecha"], "Justificación")
    for v in db.all_rows(c, f"""SELECT empleado_id, fecha_desde, fecha_hasta FROM vacaciones
                                WHERE empleado_id IN ({marcas}) AND fecha_desde <= %s AND fecha_hasta >= %s""",
                         (*ids, hasta, desde)):
        cubrir(v["empleado_id"], v["fecha_desde"], v["fecha_hasta"], "Vacaciones")
    for f in db.all_rows(c, f"SELECT empleado_id, fecha FROM francos WHERE empleado_id IN ({marcas}) AND fecha BETWEEN %s AND %s",
                         (*ids, desde, hasta)):
        cubrir(f["empleado_id"], f["fecha"], f["fecha"], "Franco")
    return salida


def _francos_banco(c, ids, desde, hasta):
    """{(empleado, fecha): minutos} de francos compensatorios del banco de horas que no fueron revertidos."""
    if not ids:
        return {}
    marcas = ",".join(["%s"] * len(ids))
    return {(int(f["empleado_id"]), _fecha(f["fecha"])): int(f["minutos"]) for f in db.all_rows(
        c, f"""SELECT m.empleado_id, m.fecha, -SUM(m.minutos) AS minutos FROM jl_banco_movimientos m
               WHERE m.tipo = 'franco' AND m.empleado_id IN ({marcas}) AND m.fecha BETWEEN %s AND %s
                 AND NOT EXISTS (SELECT 1 FROM jl_banco_movimientos r WHERE r.revierte_id = m.id)
               GROUP BY m.empleado_id, m.fecha""", (*ids, desde, hasta))}


def _huella(marcas, fecha):
    """Huella de las marcas del dia (y del siguiente, por un turno de medianoche): cambia si alguien las corrige."""
    relevantes = sorted((str(m["id"]), str(m["fecha"]), str(m["hora"]), str(m["accion"])) for m in marcas
                        if _fecha(m["fecha"]) in (fecha, fecha + dt.timedelta(days=1)))
    return hashlib.sha1(json.dumps(relevantes).encode()).hexdigest()[:16]


CLAVES = ("programados", "fichados", "trabajados", "ordinarios", "excedente", "faltantes", "nocturnos", "descanso")


def _totales(dias):
    t = {k: sum(d[k] for d in dias) for k in CLAVES}
    cuenta = defaultdict(int)
    for d in dias:
        cuenta[d["estado"]] += 1
    t.update(dias_revisar=cuenta["revisar"], dias_ausente=cuenta["ausente"], dias_licencia=cuenta["licencia"],
             dias_trabajados=sum(1 for d in dias if d["fichados"]),
             feriados_trabajados=sum(1 for d in dias if d["tipo_dia"] == "feriado" and d["fichados"]))
    return t


def _semanas(dias, resolutor, eid):
    """Trabajado por semana (lunes a domingo) contra las horas semanales pactadas (informativo)."""
    grupos = defaultdict(list)
    for d in dias:
        grupos[d["fecha"] - dt.timedelta(days=d["fecha"].weekday())].append(d)
    salida = []
    for lunes, lista in sorted(grupos.items()):
        regla, _, _ = resolutor.efectivas(eid, max(lunes, dias[0]["fecha"]))
        pactadas = int(float(regla["horas_semanales"]) * 60) if regla.get("horas_semanales") else None
        trabajados = sum(d["trabajados"] for d in lista)
        salida.append(dict(lunes=lunes, cierre=lista[-1]["fecha"], dias=len(lista), trabajados=trabajados, pactadas=pactadas,
                           exceso=max(0, trabajados - pactadas) if pactadas and len(lista) == 7 else None))
    return salida


def _alertas(dias, pares):
    alertas = []
    for d in dias:
        if d["estado"] == "revisar":
            continue  # fichadas incompletas o salida automatica: las horas no son reales
        if d["trabajados"] > JORNADA_LARGA_MIN:
            alertas.append(dict(fecha=d["fecha"], tipo="jornada_larga", texto=f"Jornada de {horas(d['trabajados'])} h"))
        elif d["pactadas_diarias"] and d["trabajados"] > d["pactadas_diarias"]:
            alertas.append(dict(fecha=d["fecha"], tipo="exceso_diario",
                                texto=f"{horas(d['trabajados'])} h trabajadas; pactadas {horas(d['pactadas_diarias'])} h"))
    periodo = {d["fecha"] for d in dias}
    fechas = sorted(f for f, lista in pares.items() if lista and (f in periodo or f + dt.timedelta(days=1) in periodo))
    for anterior, siguiente in zip(fechas, fechas[1:]):
        fin = max(f for _, f, _ in pares[anterior])
        inicio = min(i for i, _, _ in pares[siguiente])
        descanso = int((inicio - fin).total_seconds() // 60)
        if siguiente in periodo and 0 <= descanso < DESCANSO_ENTRE_JORNADAS_MIN:
            alertas.append(dict(fecha=siguiente, tipo="descanso_corto",
                                texto=f"Solo {horas(descanso)} h de descanso desde la jornada anterior (mínimo 12 h)"))
    return sorted(alertas, key=lambda a: a["fecha"])


# --- Calculos registrados ---------------------------------------------------------------------------------------

def registrar(who, desde, hasta, filtros=None, motivo=None):
    """Guarda una foto del calculo del periodo. No borra los anteriores."""
    datos = calcular(who["empresa_id"], desde, hasta, filtros)
    filas = [(r["empleado"]["id"], d) for r in datos["empleados"] for d in r["dias"]]
    with db.transaction() as c:
        calculo_id = db.insert(c, "jl_calculos", dict(
            empresa_id=who["empresa_id"], desde=datos["desde"], hasta=datos["hasta"], empleados=len(datos["empleados"]),
            dias=len(filas), motivo=(str(motivo or "").strip() or None), creado_por=who.get("user_id"), creado_at=now_local()))
        for i in range(0, len(filas), 500):
            c.executemany(
                """INSERT INTO jl_resultados_dia
                   (calculo_id, empresa_id, empleado_id, fecha, estado, tipo_dia, minutos_programados, minutos_fichados,
                    minutos_trabajados, minutos_ordinarios, minutos_excedente, minutos_faltantes, minutos_nocturnos,
                    minutos_descanso, cruza_medianoche, reglas, marcas_huella, detalle)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                [(calculo_id, who["empresa_id"], eid, d["fecha"], d["estado"], d["tipo_dia"], d["programados"], d["fichados"],
                  d["trabajados"], d["ordinarios"], d["excedente"], d["faltantes"], d["nocturnos"], d["descanso"],
                  int(d["cruza_medianoche"]), ",".join(str(x) for x in d["reglas"])[:255] or None, d["huella"],
                  json.dumps(dict(pares=d["pares"], sueltas=d["sueltas"], bloques=d["bloques"], motivo=d["motivo"],
                                  horario=d["horario"], feriado=d["feriado"]), ensure_ascii=False))
                 for eid, d in filas[i:i + 500]])
    return calculo_id, datos


def ultimo_calculo(empresa_id, desde, hasta):
    """El ultimo calculo registrado que cubre el periodo (o None) y, por empleado y dia, la huella registrada."""
    with db.transaction(read_only=True) as c:
        calculo = db.one(c, """SELECT k.*, u.usuario FROM jl_calculos k LEFT JOIN usuarios u ON u.id = k.creado_por
                               WHERE k.empresa_id = %s AND k.desde <= %s AND k.hasta >= %s
                               ORDER BY k.id DESC LIMIT 1""", (empresa_id, desde, hasta))
        if not calculo:
            return None, {}
        huellas = {(int(f["empleado_id"]), _fecha(f["fecha"])): f["marcas_huella"] for f in db.all_rows(
            c, "SELECT empleado_id, fecha, marcas_huella FROM jl_resultados_dia WHERE calculo_id = %s AND fecha BETWEEN %s AND %s",
            (calculo["id"], desde, hasta))}
    return calculo, huellas


def calculos(empresa_id, limite=20):
    with db.transaction(read_only=True) as c:
        return db.all_rows(c, """SELECT k.*, u.usuario FROM jl_calculos k LEFT JOIN usuarios u ON u.id = k.creado_por
                                 WHERE k.empresa_id = %s ORDER BY k.id DESC LIMIT %s""", (empresa_id, limite))


# --- Exportacion ------------------------------------------------------------------------------------------------

def exportar_xlsx(datos, empresa_nombre=""):
    """Excel con el resumen por empleado y el detalle diario. Horas en h:mm (texto) y en horas decimales."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    titulo = Font(bold=True, size=12)
    cabecera, relleno = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597")
    periodo = f"{datos['desde']:%d/%m/%Y} al {datos['hasta']:%d/%m/%Y}"

    def hoja(ws, nombre, columnas, filas):
        ws.title = nombre
        ws["A1"] = f"Horas trabajadas · {empresa_nombre} · {periodo}".replace(" ·  ·", " ·")
        ws["A1"].font = titulo
        ws["A2"] = (f"Calculado el {datos['calculado_at']:%d/%m/%Y %H:%M} desde los fichajes. El excedente no es hora "
                    "extra autorizada: su clasificación y autorización es la etapa de horas extras.")
        for col, texto in enumerate(columnas, 1):
            celda = ws.cell(row=4, column=col, value=texto)
            celda.font, celda.fill = cabecera, relleno
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for fila in filas:
            ws.append(list(fila))
        ws.freeze_panes = "B5"
        for col in range(1, len(columnas) + 1):
            largo = max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(4, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(10, largo + 2), 48)

    def h(m):
        return round((m or 0) / 60, 2)

    libro = Workbook()
    hoja(libro.active, "Resumen", [
        "Legajo", "Empleado", "Sucursal", "Sector", "Programadas", "Trabajadas", "Ordinarias", "Excedente", "Faltantes",
        "Nocturnas", "Programadas (h)", "Trabajadas (h)", "Excedente (h)", "Faltantes (h)", "Nocturnas (h)",
        "Días trabajados", "Días a revisar", "Ausencias", "Licencias", "Feriados trabajados"],
        [(r["empleado"]["legajo"], f"{r['empleado']['apellido']}, {r['empleado']['nombre']}", r["empleado"]["sucursal_nombre"],
          r["empleado"]["sector_nombre"], *(horas(r["totales"][k]) for k in ("programados", "trabajados", "ordinarios",
                                                                           "excedente", "faltantes", "nocturnos")),
          *(h(r["totales"][k]) for k in ("programados", "trabajados", "excedente", "faltantes", "nocturnos")),
          r["totales"]["dias_trabajados"], r["totales"]["dias_revisar"], r["totales"]["dias_ausente"],
          r["totales"]["dias_licencia"], r["totales"]["feriados_trabajados"]) for r in datos["empleados"]])
    hoja(libro.create_sheet(), "Detalle diario", [
        "Legajo", "Empleado", "Fecha", "Día", "Tipo de día", "Estado", "Horario", "Fichadas", "Programadas", "Trabajadas",
        "Excedente", "Faltantes", "Nocturnas", "Descanso", "Cruza medianoche", "Observación"],
        [(r["empleado"]["legajo"], f"{r['empleado']['apellido']}, {r['empleado']['nombre']}", d["fecha"],
          DIAS_CORTOS[d["fecha"].weekday()], TIPOS_DIA[d["tipo_dia"]], ESTADOS[d["estado"]],
          " ".join(d["bloques"]) or (d["horario"] or ""), " ".join(d["pares"] + d["sueltas"]),
          *(horas(d[k]) for k in ("programados", "trabajados", "excedente", "faltantes", "nocturnos", "descanso")),
          "Sí" if d["cruza_medianoche"] else "", d["motivo"] or d["feriado"] or d["licencia"] or "")
         for r in datos["empleados"] for d in r["dias"]])
    for celda in libro["Detalle diario"]["C"][4:]:
        celda.number_format = "DD/MM/YYYY"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
