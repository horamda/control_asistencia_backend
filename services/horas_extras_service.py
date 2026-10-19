"""Horas extras (jornada laboral, etapa C): autorizacion en dos pasos, cierre del periodo mensual y liquidacion.

- El excedente que calcula el motor (etapa B) nunca se vuelve hora extra solo: alguien lo decide. El motor propone el
  reparto entre recargo habil y de descanso (LCT art. 201) y quien decide puede aprobar menos minutos o rechazar.
- Dos pasos: el jefe (usuario del panel vinculado a un empleado, con su cadena de reportes) da el visto bueno o
  rechaza; RR. HH. (rol admin o rrhh) decide. RR. HH. puede decidir sin visto bueno: queda registrado.
- Sin fila en jl_horas_extra = pendiente (sale en vivo del calculo). Cada paso agrega un evento que no se edita.
- Topes del convenio (mes y anio): superarlos exige confirmar y escribir el motivo; queda marcado.
- Cierre mensual: exige el mes terminado, sin dias a revisar, sin excedentes sin decision de RR. HH. y sin decisiones
  que superen lo que hoy marcan los fichajes. Registra el calculo, pasa lo aprobado a "liquidada" y congela el mes.
  Reabrir pide motivo. La liquidacion sale de la foto del cierre."""

import calendar
import datetime as dt
from collections import defaultdict

from repositories import carga_repository as db
from services import banco_horas_service as banco
from services import jornada_calculo_service as calc
from services.carga_service import CargaError as Error, now_local
from services.vacaciones_semanas_service import equipo

ESTADOS = {
    "pendiente": "Pendiente del jefe",
    "visto_jefe": "Con visto bueno del jefe",
    "rechazo_jefe": "Rechazada por el jefe",
    "aprobada": "Aprobada",
    "rechazada": "Rechazada",
    "liquidada": "Liquidada",
    "compensada": "Compensada en el banco",
}
ACCIONES = {
    "visto_bueno": "Visto bueno del jefe",
    "rechazo_jefe": "Rechazo del jefe",
    "aprobar": "Aprobación de RR. HH.",
    "rechazar": "Rechazo de RR. HH.",
    "liquidar": "Liquidada al cerrar el mes",
    "compensar": "Acreditada al banco al cerrar el mes",
    "reabrir": "Mes reabierto",
}


# --- Utilidades -------------------------------------------------------------------------------------------------

def mes_de(valor):
    """'AAAA-MM' (o una fecha) -> (primer dia, ultimo dia)."""
    try:
        anio, mes = (int(x) for x in str(valor or "")[:7].split("-"))
        primero = dt.date(anio, mes, 1)
    except ValueError:
        raise Error("Mes inválido (use AAAA-MM).") from None
    return primero, primero.replace(day=calendar.monthrange(anio, mes)[1])


def _fecha(valor):
    try:
        return valor if isinstance(valor, dt.date) else dt.date.fromisoformat(str(valor)[:10])
    except ValueError:
        raise Error("Fecha inválida.") from None


def _minutos(valor, nombre):
    texto = str(valor if valor is not None else "").strip()
    if texto == "":
        return 0
    if ":" in texto:  # admite H:MM
        h, _, m = texto.partition(":")
        if h.isdigit() and m.isdigit() and int(m) < 60:
            return int(h) * 60 + int(m)
    elif texto.isdigit():
        return int(texto)
    raise Error(f"{nombre}: use minutos o H:MM.")


def estado_visible(fila):
    """Estado de una decision para mostrar (el rechazo del jefe espera la decision de RR. HH.)."""
    if not fila:
        return "pendiente"
    if fila["estado"] == "rechazada" and not fila["resuelto_por"]:
        return "rechazo_jefe"
    return fila["estado"]


def _decidida_por_rrhh(fila):
    return bool(fila) and fila["estado"] in ("aprobada", "rechazada", "liquidada", "compensada") and bool(fila["resuelto_por"])


def _aprobados(fila):
    return (fila["minutos_habil"] + fila["minutos_descanso"]) if fila and fila["estado"] in ("aprobada", "liquidada", "compensada") else 0


# --- Alcance: RR. HH. ve todo; un jefe, su cadena -----------------------------------------------------------------

def alcance(actor):
    """None = toda la empresa (RR. HH.); si no, el conjunto de empleados de la cadena del jefe (puede ser vacio)."""
    if actor.get("global"):
        return None
    if not actor.get("empleado_id"):
        return set()
    with db.transaction(read_only=True) as c:
        yo = db.one(c, "SELECT id, empresa_id, reporta_a_empleado_id FROM empleados WHERE id = %s AND empresa_id = %s",
                    (int(actor["empleado_id"]), actor["empresa_id"]))
        return set(equipo(c, yo)["visibles"]) if yo else set()


def _periodo(c, empresa_id, primero, lock=False):
    return db.one(c, "SELECT * FROM jl_periodos WHERE empresa_id = %s AND mes = %s" + (" FOR UPDATE" if lock else ""),
                  (empresa_id, primero))


def _decisiones(c, empresa_id, desde, hasta, ids=None):
    filtro, args = "", []
    if ids is not None:
        if not ids:
            return {}
        filtro = f" AND empleado_id IN ({','.join(['%s'] * len(ids))})"
        args = sorted(ids)
    return {(int(f["empleado_id"]), f["fecha"]): f for f in db.all_rows(
        c, f"SELECT * FROM jl_horas_extra WHERE empresa_id = %s AND fecha BETWEEN %s AND %s{filtro}",
        (empresa_id, desde, hasta, *args))}


def _acumulado(c, empleado_id, fecha):
    """Minutos aprobados o liquidados del empleado en el mes y en el anio de la fecha, sin contar ese dia."""
    fila = db.one(c, """SELECT
                          COALESCE(SUM(CASE WHEN fecha BETWEEN %s AND %s THEN minutos_habil + minutos_descanso END), 0) AS mes,
                          COALESCE(SUM(minutos_habil + minutos_descanso), 0) AS anio
                        FROM jl_horas_extra WHERE empleado_id = %s AND fecha BETWEEN %s AND %s AND fecha <> %s
                          AND estado IN ('aprobada', 'liquidada', 'compensada')""",
                  (*mes_de(fecha), int(empleado_id), dt.date(fecha.year, 1, 1), dt.date(fecha.year, 12, 31), fecha))
    return int(fila["mes"]), int(fila["anio"])


# --- Bandeja ----------------------------------------------------------------------------------------------------

def bandeja(actor, mes, filtros=None):
    """Dias con excedente (o con decision) del mes, en el alcance del actor, con su estado y lo que puede hacer."""
    desde, hasta = mes_de(mes)
    visibles = alcance(actor)
    filtros = dict(filtros or {})
    if visibles is not None:
        filtros["ids"] = visibles
    datos = calc.calcular(actor["empresa_id"], desde, hasta, filtros)
    ids = {int(r["empleado"]["id"]) for r in datos["empleados"]}
    with db.transaction(read_only=True) as c:
        decisiones = _decisiones(c, actor["empresa_id"], desde, hasta, ids)
        periodo = _periodo(c, actor["empresa_id"], desde)
        acumulados = {int(f["empleado_id"]): (int(f["mes"]), int(f["anio"])) for f in db.all_rows(
            c, f"""SELECT empleado_id, COALESCE(SUM(minutos_habil + minutos_descanso), 0) AS anio,
                          COALESCE(SUM(CASE WHEN fecha BETWEEN %s AND %s THEN minutos_habil + minutos_descanso END), 0) AS mes
                   FROM jl_horas_extra WHERE empleado_id IN ({','.join(['%s'] * len(ids))}) AND fecha BETWEEN %s AND %s
                     AND estado IN ('aprobada', 'liquidada', 'compensada') GROUP BY empleado_id""",
            (desde, hasta, *sorted(ids), dt.date(desde.year, 1, 1), dt.date(desde.year, 12, 31)))} if ids else {}
    cerrado = bool(periodo and periodo["estado"] == "cerrado")
    items, revisar = [], []
    for r in datos["empleados"]:
        eid = int(r["empleado"]["id"])
        for d in r["dias"]:
            fila = decisiones.get((eid, d["fecha"]))
            if d["estado"] == "revisar":
                revisar.append(dict(empleado=r["empleado"], dia=d))
            if not d["excedente"] and not fila:
                continue
            items.append(_item(actor, r["empleado"], d, fila, acumulados.get(eid, (0, 0)), cerrado))
    return dict(desde=desde, hasta=hasta, items=items, revisar=revisar, periodo=periodo, cerrado=cerrado,
                alcance="empresa" if visibles is None else "equipo", calculado_at=datos["calculado_at"])


def _item(actor, empleado, dia, fila, acumulado, cerrado):
    estado = estado_visible(fila)
    aprobados = _aprobados(fila)
    tope_mes, tope_anio = dia["topes"]
    propios_mes, propios_anio = acumulado[0] - aprobados, acumulado[1] - aprobados
    habil, descanso = dia["propuesta"]
    supera = ((tope_mes is not None and propios_mes + habil + descanso > tope_mes) or
              (tope_anio is not None and propios_anio + habil + descanso > tope_anio))
    return dict(
        empleado=empleado, dia=dia, fila=fila, estado=estado, propuesta=dia["propuesta"], acumulado=acumulado,
        supera_tope=supera, aprobados=aprobados,
        cambio=bool(fila) and fila["marcas_huella"] != dia["huella"] and fila["estado"] not in ("liquidada", "compensada"),
        excede=aprobados > dia["excedente"],
        puede=puede(actor, empleado, dia, fila, cerrado),
    )


def puede(actor, empleado, dia, fila, cerrado):
    """Acciones disponibles para el actor sobre ese dia: subconjunto de ('jefe', 'rrhh')."""
    if cerrado or not actor.get("aprobar") or int(empleado["id"]) == int(actor.get("empleado_id") or 0):
        return ()
    if fila and fila["estado"] in ("liquidada", "compensada"):
        return ()
    if actor.get("global"):
        return ("rrhh",) if dia["estado"] == "normal" and dia["excedente"] or fila else ()
    if fila and fila["resuelto_por"]:
        return ()  # RR. HH. ya decidio
    return ("jefe",) if dia["estado"] == "normal" and dia["excedente"] else ()


# --- Decidir ----------------------------------------------------------------------------------------------------

def _dia_actual(actor, empleado_id, fecha):
    visibles = alcance(actor)
    if visibles is not None and int(empleado_id) not in visibles:
        raise Error("El empleado no está en su equipo.", 403)
    datos = calc.calcular(actor["empresa_id"], fecha, fecha, {"empleado_id": empleado_id})
    if not datos["empleados"]:
        raise Error("Empleado no encontrado.", 404)
    return datos["empleados"][0]["empleado"], datos["empleados"][0]["dias"][0]


def detalle(actor, empleado_id, fecha):
    fecha = _fecha(fecha)
    empleado, dia = _dia_actual(actor, empleado_id, fecha)
    with db.transaction(read_only=True) as c:
        fila = db.one(c, "SELECT * FROM jl_horas_extra WHERE empleado_id = %s AND fecha = %s", (int(empleado_id), fecha))
        eventos = db.all_rows(c, """SELECT e.*, u.usuario FROM jl_horas_extra_eventos e LEFT JOIN usuarios u ON u.id = e.usuario_id
                                    WHERE e.horas_extra_id = %s ORDER BY e.id""", (fila["id"],)) if fila else []
        periodo = _periodo(c, actor["empresa_id"], fecha.replace(day=1))
        acumulado = _acumulado(c, empleado_id, fecha)
        acuerdo = banco.acuerdo_vigente(c, empleado_id, fecha)
    aprobados = _aprobados(fila)
    acumulado = (acumulado[0] + aprobados, acumulado[1] + aprobados)
    item = _item(actor, empleado, dia, fila, acumulado, bool(periodo and periodo["estado"] == "cerrado"))
    return dict(item, eventos=eventos, periodo=periodo, acuerdo_banco=acuerdo)


def decidir(actor, empleado_id, fecha, accion, minutos_habil=None, minutos_descanso=None, motivo=None, confirmar_tope=False,
            destino="pago"):
    """accion: visto_bueno o rechazo_jefe (jefe), aprobar o rechazar (RR. HH.). Sin minutos, se toma la propuesta.
    destino (al aprobar): 'pago' o 'banco' (banco de horas, si tiene acuerdo vigente)."""
    fecha = _fecha(fecha)
    empleado, dia = _dia_actual(actor, empleado_id, fecha)
    return _decidir(actor, empleado, dia, accion, minutos_habil, minutos_descanso, motivo, confirmar_tope, destino)


def decidir_lote(actor, mes, claves, accion):
    """Aplica la propuesta del motor a varios dias [(empleado_id, fecha)]. Devuelve (hechos, errores)."""
    if accion not in ("visto_bueno", "aprobar"):
        raise Error("En lote solo se puede dar el visto bueno o aprobar la propuesta.")
    desde, hasta = mes_de(mes)
    claves = {(int(e), _fecha(f)) for e, f in claves}
    if not claves:
        raise Error("Elegí al menos un día.")
    visibles = alcance(actor)
    ids = {e for e, _ in claves} if visibles is None else {e for e, _ in claves} & visibles
    datos = calc.calcular(actor["empresa_id"], desde, hasta, {"ids": ids})
    dias = {(int(r["empleado"]["id"]), d["fecha"]): (r["empleado"], d) for r in datos["empleados"] for d in r["dias"]}
    hechos, errores = 0, []
    for clave in sorted(claves):
        if clave not in dias:
            errores.append(dict(clave=clave, error="fuera de su equipo o del mes"))
            continue
        try:
            _decidir(actor, *dias[clave], accion)
            hechos += 1
        except Error as error:
            errores.append(dict(clave=clave, nombre=f"{dias[clave][0]['apellido']}, {dias[clave][0]['nombre']}", error=str(error)))
    return hechos, errores


def _decidir(actor, empleado, dia, accion, minutos_habil=None, minutos_descanso=None, motivo=None, confirmar_tope=False,
             destino="pago"):
    if accion not in ("visto_bueno", "rechazo_jefe", "aprobar", "rechazar"):
        raise Error("Acción inválida.")
    if destino not in ("pago", "banco"):
        raise Error("Destino inválido.")
    destino = destino if accion == "aprobar" else "pago"
    eid, fecha = int(empleado["id"]), dia["fecha"]
    motivo = (str(motivo or "").strip() or None)
    rechazo = accion in ("rechazo_jefe", "rechazar")
    if rechazo:
        habil = descanso = 0
    elif minutos_habil is None and minutos_descanso is None:
        habil, descanso = dia["propuesta"]
    else:
        habil, descanso = _minutos(minutos_habil, "Minutos hábiles"), _minutos(minutos_descanso, "Minutos de descanso")
    if dia["estado"] != "normal" or not dia["excedente"]:
        if not rechazo:
            raise Error("Ese día no tiene excedente para aprobar (o tiene fichadas a revisar).")
    if not rechazo and not habil + descanso:
        raise Error("Indicá los minutos a aprobar o rechazá.")
    if habil + descanso > dia["excedente"]:
        raise Error(f"No se pueden aprobar más minutos que el excedente del día ({calc.horas(dia['excedente'])} h).")
    if rechazo and not motivo:
        raise Error("Indicá el motivo del rechazo.")
    ahora = now_local().replace(microsecond=0)
    with db.transaction() as c:
        periodo = _periodo(c, actor["empresa_id"], fecha.replace(day=1), lock=True)
        if periodo and periodo["estado"] == "cerrado":
            raise Error("El mes está cerrado: para cambiarlo hay que reabrirlo.")
        fila = db.one(c, "SELECT * FROM jl_horas_extra WHERE empleado_id = %s AND fecha = %s FOR UPDATE", (eid, fecha))
        permitido = puede(actor, empleado, dia, fila, False)
        if ("jefe" if accion in ("visto_bueno", "rechazo_jefe") else "rrhh") not in permitido:
            raise Error("No puede tomar esa decisión sobre este día.", 403)
        if destino == "banco":
            banco.validar_destino_banco(c, eid, fecha, habil, descanso, fila and fila["id"])
        supera = False
        if not rechazo:
            mes, anio = _acumulado(c, eid, fecha)
            tope_mes, tope_anio = dia["topes"]
            supera = ((tope_mes is not None and mes + habil + descanso > tope_mes) or
                      (tope_anio is not None and anio + habil + descanso > tope_anio))
            if supera and accion == "aprobar" and not (confirmar_tope and motivo):
                topes = " y ".join(t for t in (f"{tope_mes // 60} h por mes" if tope_mes is not None else "",
                                               f"{tope_anio // 60} h por año" if tope_anio is not None else "") if t)
                raise Error(f"Supera el tope de horas extras del convenio ({topes}). Para aprobar igual, confirmá y "
                            "escribí el motivo.")
        estado = {"visto_bueno": "visto_jefe", "rechazo_jefe": "rechazada", "aprobar": "aprobada", "rechazar": "rechazada"}[accion]
        valores = dict(estado=estado, minutos_excedente=dia["excedente"], propuesta_habil=dia["propuesta"][0],
                       propuesta_descanso=dia["propuesta"][1], minutos_habil=habil, minutos_descanso=descanso,
                       recargo_habil=dia["recargos"][0], recargo_descanso=dia["recargos"][1], supera_tope=int(supera),
                       marcas_huella=dia["huella"], motivo=motivo, destino=destino, actualizado_at=ahora)
        if accion in ("visto_bueno", "rechazo_jefe"):
            valores.update(jefe_por=actor["user_id"], jefe_at=ahora, resuelto_por=None, resuelto_at=None)
        else:
            valores.update(resuelto_por=actor["user_id"], resuelto_at=ahora)
        if fila:
            c.execute(f"UPDATE jl_horas_extra SET {', '.join(f'{k} = %s' for k in valores)} WHERE id = %s",
                      (*valores.values(), fila["id"]))
            hid = fila["id"]
        else:
            hid = db.insert(c, "jl_horas_extra", dict(valores, empresa_id=actor["empresa_id"], empleado_id=eid, fecha=fecha,
                                                      creado_at=ahora))
        db.insert(c, "jl_horas_extra_eventos", dict(
            horas_extra_id=hid, empresa_id=actor["empresa_id"], accion=accion, estado_anterior=fila and fila["estado"],
            estado_nuevo=estado, minutos_habil=habil, minutos_descanso=descanso, supera_tope=int(supera), motivo=motivo,
            usuario_id=actor["user_id"], creado_at=ahora))
    return hid


# --- Cierre del periodo -----------------------------------------------------------------------------------------

def estado_cierre(actor, mes):
    """Estado del mes y lo que falta para cerrarlo (toda la empresa)."""
    desde, hasta = mes_de(mes)
    datos = calc.calcular(actor["empresa_id"], desde, hasta)
    with db.transaction(read_only=True) as c:
        decisiones = _decisiones(c, actor["empresa_id"], desde, hasta)
        periodo = _periodo(c, actor["empresa_id"], desde)
    pendientes, revisar, excedidas = [], [], []
    resumen = defaultdict(int)
    for r in datos["empleados"]:
        eid = int(r["empleado"]["id"])
        for d in r["dias"]:
            fila = decisiones.get((eid, d["fecha"]))
            item = dict(empleado=r["empleado"], dia=d, fila=fila, estado=estado_visible(fila))
            if d["estado"] == "revisar":
                revisar.append(item)
            if d["excedente"] and not _decidida_por_rrhh(fila):
                pendientes.append(item)
            if fila and fila["estado"] == "aprobada" and _aprobados(fila) > d["excedente"]:
                excedidas.append(item)
            if fila:
                resumen[fila["estado"]] += 1
                resumen["minutos_" + fila["estado"]] += _aprobados(fila)
    terminado = hasta < calc._hoy()
    return dict(desde=desde, hasta=hasta, periodo=periodo, pendientes=pendientes, revisar=revisar, excedidas=excedidas,
                terminado=terminado, resumen=dict(resumen), empleados=len(datos["empleados"]),
                cerrado=bool(periodo and periodo["estado"] == "cerrado"),
                se_puede_cerrar=terminado and not (pendientes or revisar or excedidas) and not (periodo and periodo["estado"] == "cerrado"))


def cerrar(actor, mes):
    if not actor.get("global"):
        raise Error("Solo RR. HH. puede cerrar el mes.", 403)
    estado = estado_cierre(actor, mes)
    if estado["cerrado"]:
        raise Error("El mes ya está cerrado.")
    if not estado["terminado"]:
        raise Error("El mes todavía no terminó.")
    faltan = [f"{len(lista)} {texto}" for lista, texto in (
        (estado["pendientes"], "días con excedente sin decisión de RR. HH."),
        (estado["revisar"], "días con fichadas a revisar"),
        (estado["excedidas"], "aprobaciones que superan lo que hoy marcan los fichajes")) if lista]
    if faltan:
        raise Error("No se puede cerrar: " + "; ".join(faltan) + ".")
    desde = estado["desde"]
    calculo_id, _ = calc.registrar(actor, desde, estado["hasta"], motivo=f"Cierre de {desde:%m/%Y}")
    ahora = now_local().replace(microsecond=0)
    with db.transaction() as c:
        periodo = _periodo(c, actor["empresa_id"], desde, lock=True)
        if periodo and periodo["estado"] == "cerrado":
            raise Error("El mes ya está cerrado.")
        datos = dict(estado="cerrado", calculo_id=calculo_id, cerrado_por=actor["user_id"], cerrado_at=ahora)
        if periodo:
            c.execute("UPDATE jl_periodos SET estado=%s, calculo_id=%s, cerrado_por=%s, cerrado_at=%s WHERE id=%s",
                      (*datos.values(), periodo["id"]))
            pid = periodo["id"]
        else:
            pid = db.insert(c, "jl_periodos", dict(datos, empresa_id=actor["empresa_id"], mes=desde))
        filas = db.all_rows(c, """SELECT * FROM jl_horas_extra WHERE empresa_id = %s AND fecha BETWEEN %s AND %s
                                  AND estado = 'aprobada' AND destino = 'pago' FOR UPDATE""",
                            (actor["empresa_id"], desde, estado["hasta"]))
        for f in filas:
            c.execute("UPDATE jl_horas_extra SET estado='liquidada', periodo_id=%s, actualizado_at=%s WHERE id=%s",
                      (pid, ahora, f["id"]))
            db.insert(c, "jl_horas_extra_eventos", dict(
                horas_extra_id=f["id"], empresa_id=actor["empresa_id"], accion="liquidar", estado_anterior="aprobada",
                estado_nuevo="liquidada", minutos_habil=f["minutos_habil"], minutos_descanso=f["minutos_descanso"],
                supera_tope=f["supera_tope"], motivo=f"Cierre de {desde:%m/%Y}", usuario_id=actor["user_id"], creado_at=ahora))
        del_banco = banco.al_cerrar(c, actor, pid, desde, estado["hasta"], calculo_id)
    return pid, len(filas), del_banco


def reabrir(actor, mes, motivo):
    if not actor.get("global"):
        raise Error("Solo RR. HH. puede reabrir el mes.", 403)
    motivo = str(motivo or "").strip()[:255]
    if not motivo:
        raise Error("Indicá el motivo de la reapertura.")
    desde, hasta = mes_de(mes)
    ahora = now_local().replace(microsecond=0)
    with db.transaction() as c:
        periodo = _periodo(c, actor["empresa_id"], desde, lock=True)
        if not periodo or periodo["estado"] != "cerrado":
            raise Error("El mes no está cerrado.")
        c.execute("UPDATE jl_periodos SET estado='reabierto', reabierto_por=%s, reabierto_at=%s, motivo_reapertura=%s WHERE id=%s",
                  (actor["user_id"], ahora, motivo, periodo["id"]))
        filas = db.all_rows(c, "SELECT * FROM jl_horas_extra WHERE periodo_id = %s AND estado = 'liquidada' FOR UPDATE",
                            (periodo["id"],))
        for f in filas:
            c.execute("UPDATE jl_horas_extra SET estado='aprobada', periodo_id=NULL, actualizado_at=%s WHERE id=%s", (ahora, f["id"]))
            db.insert(c, "jl_horas_extra_eventos", dict(
                horas_extra_id=f["id"], empresa_id=actor["empresa_id"], accion="reabrir", estado_anterior="liquidada",
                estado_nuevo="aprobada", minutos_habil=f["minutos_habil"], minutos_descanso=f["minutos_descanso"],
                supera_tope=f["supera_tope"], motivo=motivo, usuario_id=actor["user_id"], creado_at=ahora))
        banco.al_reabrir(c, actor, periodo["id"], motivo)
    return periodo["id"], len(filas)


def periodos(empresa_id, limite=12):
    with db.transaction(read_only=True) as c:
        return db.all_rows(c, """SELECT p.*, uc.usuario AS cerrado_usuario, ur.usuario AS reabierto_usuario FROM jl_periodos p
                                 LEFT JOIN usuarios uc ON uc.id = p.cerrado_por LEFT JOIN usuarios ur ON ur.id = p.reabierto_por
                                 WHERE p.empresa_id = %s ORDER BY p.mes DESC LIMIT %s""", (empresa_id, limite))


# --- Liquidacion (desde la foto del cierre) ------------------------------------------------------------------------

def liquidacion(actor, mes):
    if not actor.get("global"):
        raise Error("Solo RR. HH. puede exportar la liquidación.", 403)
    desde, hasta = mes_de(mes)
    with db.transaction(read_only=True) as c:
        periodo = _periodo(c, actor["empresa_id"], desde)
        if not periodo or periodo["estado"] != "cerrado":
            raise Error("La liquidación sale de un mes cerrado: cerralo primero.")
        filas = db.all_rows(c, """SELECT e.id, e.legajo, e.apellido, e.nombre, e.dni, su.nombre AS sucursal_nombre,
                                         SUM(r.minutos_trabajados) AS trabajados, SUM(r.minutos_nocturnos) AS nocturnos,
                                         SUM(r.estado = 'ausente') AS ausencias, SUM(r.estado = 'licencia') AS licencias,
                                         SUM(r.tipo_dia = 'feriado' AND r.minutos_fichados > 0) AS feriados
                                  FROM jl_resultados_dia r JOIN empleados e ON e.id = r.empleado_id
                                  LEFT JOIN sucursales su ON su.id = e.sucursal_id
                                  WHERE r.calculo_id = %s GROUP BY e.id, e.legajo, e.apellido, e.nombre, e.dni, su.nombre
                                  ORDER BY e.apellido, e.nombre""", (periodo["calculo_id"],))
        extras = db.all_rows(c, """SELECT h.*, e.legajo, e.apellido, e.nombre FROM jl_horas_extra h
                                   JOIN empleados e ON e.id = h.empleado_id
                                   WHERE h.periodo_id = %s AND h.estado IN ('liquidada', 'compensada')
                                   ORDER BY e.apellido, e.nombre, h.fecha""",
                             (periodo["id"],))
    por_empleado = defaultdict(lambda: dict(habil=0, descanso=0, banco=0, recargos_habil=set(), recargos_descanso=set()))
    for x in extras:
        t = por_empleado[int(x["empleado_id"])]
        if x["estado"] == "compensada":  # no se pagan: van al banco de horas
            t["banco"] += x["minutos_habil"] + x["minutos_descanso"]
            continue
        t["habil"] += x["minutos_habil"]
        t["descanso"] += x["minutos_descanso"]
        if x["minutos_habil"] and x["recargo_habil"] is not None:
            t["recargos_habil"].add(int(x["recargo_habil"]))
        if x["minutos_descanso"] and x["recargo_descanso"] is not None:
            t["recargos_descanso"].add(int(x["recargo_descanso"]))
    for f in filas:
        f.update(extras=por_empleado.get(int(f["id"]), dict(habil=0, descanso=0, banco=0, recargos_habil=set(),
                                                            recargos_descanso=set())))
    return dict(desde=desde, hasta=hasta, periodo=periodo, filas=filas, extras=extras)


def exportar_liquidacion(datos):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cabecera, relleno = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F5597")
    p = datos["periodo"]

    def hoja(ws, nombre, columnas, filas):
        ws.title = nombre
        ws["A1"] = f"Liquidación de horas · {datos['desde']:%m/%Y}"
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = (f"Mes cerrado el {p['cerrado_at']:%d/%m/%Y %H:%M} (cálculo registrado N.º {p['calculo_id']}). "
                    "Horas en decimal; solo horas extras aprobadas por RR. HH.")
        for col, texto in enumerate(columnas, 1):
            celda = ws.cell(row=4, column=col, value=texto)
            celda.font, celda.fill = cabecera, relleno
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for fila in filas:
            ws.append(list(fila))
        ws.freeze_panes = "C5"
        for col in range(1, len(columnas) + 1):
            largo = max(len(str(ws.cell(row=r, column=col).value or "")) for r in range(4, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(col)].width = min(max(10, largo + 2), 40)

    def h(m):
        return round(int(m or 0) / 60, 2)

    def pct(valores):
        return " / ".join(f"{v} %" for v in sorted(valores)) or ""

    libro = Workbook()
    hoja(libro.active, "Liquidación", [
        "Legajo", "Empleado", "DNI", "Sucursal", "Horas extras hábiles", "Recargo hábil", "Horas extras descanso y feriado",
        "Recargo descanso", "Horas extras al banco (no se pagan)", "Horas nocturnas", "Feriados trabajados", "Ausencias injustificadas", "Días de licencia",
        "Horas trabajadas"],
        [(f["legajo"], f"{f['apellido']}, {f['nombre']}", f["dni"], f["sucursal_nombre"], h(f["extras"]["habil"]),
          pct(f["extras"]["recargos_habil"]), h(f["extras"]["descanso"]), pct(f["extras"]["recargos_descanso"]),
          h(f["extras"]["banco"]), h(f["nocturnos"]), int(f["feriados"] or 0), int(f["ausencias"] or 0), int(f["licencias"] or 0), h(f["trabajados"]))
         for f in datos["filas"]])
    hoja(libro.create_sheet(), "Horas extras por día", [
        "Legajo", "Empleado", "Fecha", "Excedente", "Hábiles", "Recargo hábil", "Descanso y feriado", "Recargo descanso",
        "Destino", "Supera tope", "Motivo"],
        [(x["legajo"], f"{x['apellido']}, {x['nombre']}", x["fecha"], h(x["minutos_excedente"]), h(x["minutos_habil"]),
          f"{x['recargo_habil']} %" if x["recargo_habil"] is not None else "", h(x["minutos_descanso"]),
          f"{x['recargo_descanso']} %" if x["recargo_descanso"] is not None else "",
          "Banco de horas" if x["estado"] == "compensada" else "Pago", "Sí" if x["supera_tope"] else "",
          x["motivo"] or "") for x in datos["extras"]])
    for celda in libro["Horas extras por día"]["C"][4:]:
        celda.number_format = "DD/MM/YYYY"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
