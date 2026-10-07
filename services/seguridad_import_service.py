"""Archivo original inmutable, filas recuperables y resolución explícita."""
import hashlib
import io
import json
import unicodedata
import zipfile
from datetime import date, datetime, time
from uuid import UUID, uuid5, NAMESPACE_URL
from openpyxl import load_workbook
from services import seguridad_service as s
from services import seguridad_excel_schema as schema


class ValidationError(s.Error):
    def __init__(self, issues):
        super().__init__('La plantilla tiene errores. Corríjalos antes de importar; no se guardó ningún registro.')
        self.issues = issues


def source_id(empresa, original):
    if schema.is_template(original):
        value=str(schema.values(original).get('ID registro') or '').strip()
        try: return str(UUID(value))
        except ValueError: return value.lower()
    canonical = []
    for index in range(max(len(original['encabezados']), len(original['celdas']))):
        h = original['encabezados'][index] if index < len(original['encabezados']) else None
        v = original['celdas'][index] if index < len(original['celdas']) else None
        if h or v is not None:
            canonical.append((key(h) or f'columna_{index}', str(v).strip() if v is not None else ''))
    digest = hashlib.sha256(s.dumps(canonical).encode()).hexdigest()
    return str(uuid5(NAMESPACE_URL, f'sh-import:{empresa}:{digest}'))


def origin_row(c, empresa, original, lock=False, records=None):
    """Verifica el vínculo con una respuesta previa, sin confiar en IDs del archivo."""
    if not schema.is_template(original):
        return None
    value = schema.values(original).get('Fila origen')
    if value in (None, ''):
        return None
    row_id = s.positive_int(value, 'Fila origen')
    row = records.get(row_id) if records is not None else s.db.one(c, '''SELECT f.* FROM sh_import_filas f
        JOIN sh_importaciones i ON i.id=f.importacion_id WHERE f.id=%s AND i.empresa_id=%s'''
        + (' FOR UPDATE' if lock else ''), (row_id, empresa))
    if not row or source_id(empresa, json.loads(row['original'])) != source_id(empresa, original):
        raise s.Error('Fila origen: el identificador no corresponde a una respuesta de esta empresa.')
    return row


def key(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD',str(value or '')) if not unicodedata.combining(c)).lower().split())


def legacy_subject(original):
    def get(label,index=0):
        positions=[i for i,h in enumerate(original['encabezados']) if key(h)==key(label)]
        return original['celdas'][positions[index]] if len(positions)>index and positions[index]<len(original['celdas']) else None
    name=get('Apellido y nombre completo de la persona a informar')
    if name: return name,get('Legajo (informado)')
    safe=key(get('Que va a notificar'))=='comportamiento seguro'
    return get('Nombre Completo de quien comete la infracción',int(safe)),get('Legajo (observado)' if safe else 'Legajo (infractor)')


def legajo_key(value):
    if isinstance(value,float) and value.is_integer(): value=int(value)
    value=str(value or '').strip()
    return str(int(value)) if value.isdigit() else value


def read_file(file):
    if not file or not (file.filename or '').lower().endswith('.xlsx'):
        raise s.Error('Seleccione el Excel de respuestas (.xlsx).')
    raw=file.read(10*1024*1024+1)
    if not raw or len(raw)>10*1024*1024: raise s.Error('Máximo 10 MB por archivo.')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(i.file_size for i in z.infolist())>50*1024*1024 or len(z.infolist())>2000:
                raise s.Error('El contenido descomprimido del archivo es demasiado grande.')
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        equivalences={}
        if 'Equivalencias' in workbook.sheetnames and 'Carga' not in workbook.sheetnames:
            sheet=workbook['Equivalencias']
            if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>200:
                raise s.Error('La hoja Equivalencias excede los límites de tamaño.')
            for number,row in enumerate(sheet.values,1):
                if number==1:
                    if [key(v) for v in row[:4]]!=['nombre como fue escrito','legajo','empleado segun lista','estado']:
                        raise s.Error('Revise los encabezados de la hoja Equivalencias.')
                    continue
                if not row[0]: continue
                name=key(row[0])
                entry=dict(nombre=str(row[2] or row[0]).strip(),legajo=row[1],estado=key(row[3]))
                if name in equivalences and (key(equivalences[name]['nombre']),legajo_key(equivalences[name]['legajo']),equivalences[name]['estado'])!=(key(entry['nombre']),legajo_key(entry['legajo']),entry['estado']):
                    raise s.Error('Hay equivalencias contradictorias para una misma persona.')
                equivalences[name]=entry
        result=[]
        for sheet in ([workbook['Carga']] if 'Carga' in workbook.sheetnames else workbook):
            if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>200:
                raise s.Error('Máximo 10.000 filas y 200 columnas por hoja.')
            headers=None
            template=False
            for number,row in enumerate(sheet.values,1):
                if number>10000 or len(row)>200: raise s.Error('Máximo 10.000 filas y 200 columnas por hoja.')
                values=[v.isoformat() if isinstance(v,(date,datetime,time)) else v for v in row]
                if not any(v not in (None,'') for v in values): continue
                normalized=[key(v) for v in values]
                if sheet.title == 'Carga' and 'id registro' in normalized and 'legajos' in normalized:
                    if values[:len(schema.HEADERS)] != schema.HEADERS and values != schema.LEGACY_HEADERS:
                        raise s.Error('No cambie los encabezados ni su orden en la hoja Carga. Descargue la plantilla nuevamente.')
                    headers=values; template=True; continue
                if 'marca temporal' in normalized and 'que va a notificar' in normalized:
                    headers=values; continue
                # La cabecera de grupos permanece en el archivo, no es una respuesta.
                if headers is None: continue
                original=dict(encabezados=headers,celdas=values)
                if template:
                    original['formato']='seguridad_v2' if headers[:len(schema.HEADERS)]==schema.HEADERS else 'seguridad_v1'
                    data=schema.values(original)
                    if data.get('Fila origen') in (None,'') and not any(data.get(h) not in (None,'') for h in schema.BUSINESS): continue
                else:
                    name,legajo=legacy_subject(original)
                    if key(name) in equivalences:
                        original['equivalencia']=dict(equivalences[key(name)])
                result.append(dict(hoja=sheet.title,fila=number,original=original))
                if len(result)>5000: raise s.Error('Máximo 5.000 respuestas por importación.')
        workbook.close()
    except s.Error: raise
    except Exception as exc: raise s.Error('No se pudo leer el Excel. Use el archivo de respuestas con sus encabezados originales.') from exc
    if not result: raise s.Error('No se encontró la cabecera del formulario ni respuestas.')
    return raw,result


def suggest(original,people,catalogs,externals=()):
    if schema.is_template(original):
        return schema.parse(original,people,catalogs,externals)
    cells=original['celdas']; headers=original['encabezados']
    def get(label,index=0):
        matches=[i for i,h in enumerate(headers) if key(h)==key(label)]
        return cells[matches[index]] if len(matches)>index and matches[index]<len(cells) else None
    kind={'comportamiento seguro':'seguro','comportamiento inseguro':'inseguro'}.get(key(get('Que va a notificar')),'')
    errors=[]
    if not kind:
        errors.append('Defina el tipo; accidente e incidente estaban agrupados.' if key(get('Que va a notificar'))=='accidente-incidente' else 'Fila sin tipo informado: revisar datos originales.')
    name,legajo=legacy_subject(original)
    equivalence=original.get('equivalencia')
    external=[]
    matches=[]
    if original.get('resolucion_empleado_id'):
        # Explicit administrative correction; preserve the source cells and UUID.
        matches=[p for p in people if p['id']==original['resolucion_empleado_id']]
    elif equivalence:
        status=equivalence['estado']
        if status=='ok':
            approved=legajo_key(equivalence['legajo'])
            if not approved or (legajo not in (None,'') and legajo_key(legajo)!=approved):
                errors.append('El legajo informado contradice la equivalencia aprobada.')
            else:
                matches=[p for p in people if legajo_key(p['legajo'])==approved]
        elif status=='externo':
            if legajo not in (None,'') and key(legajo)!='externo':
                errors.append('La persona externa tiene un legajo de empleado: revise la equivalencia.')
            elif equivalence.get('externo_id'):
                if any(p['id']==equivalence['externo_id'] and p['activo'] for p in externals):
                    external=[dict(id=equivalence['externo_id'])]
                else: errors.append('La persona externa no está activa en esta empresa.')
            else:
                external=[dict(nombre=equivalence['nombre'],empresa='')]
        else:
            errors.append('Equivalencia pendiente: '+status.upper()+'. Confirme la identidad de la persona.')
    elif legajo not in (None,''):
        if key(legajo)=='externo' and name:
            external=[dict(nombre=str(name).strip(),empresa='')]
        else:
            matches=[p for p in people if legajo_key(p['legajo'])==legajo_key(legajo)]
    else:
        matches=[p for p in people if key(name) in (key(f"{p['apellido']} {p['nombre']}"),key(f"{p['nombre']} {p['apellido']}"))]
    ids=[matches[0]['id']] if len(matches)==1 else []
    if not ids and not external: errors.append('Vincule la persona con su legajo; no hay una coincidencia única.')
    date_value=get('Fecha que se observa la condición insegura-comportamiento inseguro') if kind=='inseguro' else get('Fecha del evento',1 if kind=='seguro' else 0)
    fecha=str(date_value or '')[:10]
    try:
        parsed=date.fromisoformat(fecha)
        if parsed.year<2000 or parsed>s.now_local().date(): raise ValueError()
    except ValueError:
        fecha=''; errors.append('Corrija la fecha del evento.')
    name_cat=get('Infracción que cometió') if kind=='inseguro' else get('Comportamiento Seguro') if kind=='seguro' else None
    cats=[c for c in catalogs if c['clase']=='categoria' and c['tipo']==kind and c['activo'] and key(c['nombre'])==key(name_cat)]
    cat=cats[0]['id'] if len(cats)==1 else None
    if kind and not cat: errors.append('Seleccione o cree la categoría en Catálogos.')
    desc=get('Comentarios') if kind=='inseguro' else get('Comentarios, describir la tarea que se realizo') if kind=='seguro' else get('Descripción del evento')
    if not desc: errors.append('Complete la descripción; no se inventa a partir de la categoría.')
    value=dict(tipo=kind,persona_original=str(name or ''),fecha_evento=fecha,hora_evento=str(get('Horario del Evento') or '')[:5],
        descripcion=str(desc or ''),categoria_id=cat,involucrados=ids,
        advertido=key(get('Se advirtió a la persona que estaba cometiendo el acto?')),
        lugar=get('Lugar del evento') or '',clasificacion=get('¿Es SIF (hubo riesgo de muerte), FAI(lesión leve), MTI(lesión moderada),  MDI(lesión grave), o LTI(lesión muy grave)?') or '',
        atencion=get('Requiero de') or '',lesion=get('Tipo de Lesión') or '',parte_cuerpo=get('Parte del Cuerpo lesionado') or '')
    if external: value['externos']=external
    if kind=='inseguro' and value['advertido'] not in ('si','no'): errors.append('Indique si se advirtió a la persona.')
    if not errors:
        try: s.parse(value)
        except s.Error as exc: errors.append(str(exc))
    return value,errors


def stage(empresa,user,file):
    raw,rows=read_file(file)
    sha=hashlib.sha256(raw).hexdigest()
    with s.db.transaction() as c:
        s.db.one(c,'SELECT id FROM empresas WHERE id=%s FOR UPDATE',(empresa,))
        old=s.db.one(c,'SELECT id FROM sh_importaciones WHERE empresa_id=%s AND sha256=%s',(empresa,sha))
        if old: return old['id']
        people=s.employees(c,empresa,False)
        cats=s.db.all_rows(c,'SELECT * FROM sh_catalogos WHERE empresa_id=%s',(empresa,))
        externals=s.external_people(c,empresa)
        # An approved equivalence identifies one person throughout this workbook.
        # Do not merge existing registry names: homonyms can be different people.
        external_ids={}
        for row in rows:
            original=row['original']
            eq=original.get('equivalencia')
            if not eq or eq['estado']!='externo': continue
            _,legajo=legacy_subject(original)
            if legajo not in (None,'') and key(legajo)!='externo': continue
            identity=key(eq['nombre'])
            if identity not in external_ids:
                parsed=s.parse_externals([dict(nombre=eq['nombre'],empresa='')])[0]
                record=dict(empresa_id=empresa,nombre=parsed['nombre'],empresa='',activo=1,creado_at=s.now_local())
                external_id=s.db.insert(c,'sh_personas_externas',record)
                external_ids[identity]=external_id
                externals.append(dict(record,id=external_id))
                s.audit(c,empresa,user,'crear_externo_importacion',dict(externo_id=external_id,sha256=sha))
            eq['externo_id']=external_ids[identity]
        prepared=[]; issues=[]; seen=set()
        template_ids=[source_id(empresa,r['original']) for r in rows if schema.is_template(r['original'])]
        existing={}
        origins={}
        if template_ids:
            existing={r['envio_id']:r for r in s.db.all_rows(c,
                'SELECT envio_id,contenido_hash,origen,reportante_id FROM sh_eventos WHERE empresa_id=%s AND envio_id IN ('+
                ','.join(['%s']*len(template_ids))+')',(empresa,*template_ids))}
            origin_ids=[]
            for r in rows:
                if schema.is_template(r['original']):
                    try:
                        origin_ids.append(s.positive_int(schema.values(r['original']).get('Fila origen'),'Fila origen'))
                    except s.Error:
                        pass
            if origin_ids:
                origins={r['id']:r for r in s.db.all_rows(c,
                    'SELECT f.* FROM sh_import_filas f JOIN sh_importaciones i ON i.id=f.importacion_id WHERE i.empresa_id=%s AND f.id IN ('+
                    ','.join(['%s']*len(origin_ids))+')',(empresa,*origin_ids))}
        for row in rows:
            proposed,errors=suggest(row['original'],people,cats,externals)
            if schema.is_template(row['original']):
                identifier=proposed['envio_id']
                if identifier in seen:
                    errors.append('ID registro: repetido dentro del archivo; use una fila por evento.')
                seen.add(identifier)
                try:
                    origin_row(c,empresa,row['original'],records=origins)
                except s.Error as exc:
                    errors.append(str(exc))
                old_event=existing.get(identifier)
                if old_event and not errors:
                    fingerprint=hashlib.sha256(s.dumps([s.parse(proposed),[]]).encode()).hexdigest()
                    if old_event['contenido_hash']!=fingerprint or old_event['origen']!='importacion' or old_event['reportante_id'] is not None:
                        errors.append('ID registro: ya existe con otros datos. Edite el evento en el panel; no cambie el ID para duplicarlo.')
                if errors:
                    issues.append(dict(hoja=row['hoja'],fila=row['fila'],errores=errors))
            prepared.append((row,proposed,errors))
        if issues: raise ValidationError(issues)
        batch=s.db.insert(c,'sh_importaciones',dict(empresa_id=empresa,usuario_id=user,nombre=file.filename[:255],archivo=raw,sha256=sha,creado_at=s.now_local()))
        c.executemany('INSERT INTO sh_import_filas (importacion_id,hoja,fila,original,propuesta,errores) VALUES (%s,%s,%s,%s,%s,%s)',
                      [(batch,row['hoja'],row['fila'],s.dumps(row['original']),s.dumps(proposed),s.dumps(errors)) for row,proposed,errors in prepared])
        s.audit(c,empresa,user,'importar_archivo',dict(importacion_id=batch,filas=len(rows),sha256=sha))
        return batch


def batch(empresa,batch_id,page=1):
    with s.db.transaction() as c:
        row=s.db.one(c,'SELECT id,nombre,creado_at FROM sh_importaciones WHERE id=%s AND empresa_id=%s',(batch_id,empresa))
        if not row: raise s.Error('Importación no encontrada.',404)
        row['totales']=s.db.one(c,'SELECT COUNT(*) total,SUM(evento_id IS NOT NULL) resueltas FROM sh_import_filas WHERE importacion_id=%s',(batch_id,))
        row['filas']=s.db.all_rows(c,'SELECT id,hoja,fila,propuesta,errores,evento_id FROM sh_import_filas WHERE importacion_id=%s ORDER BY fila,id LIMIT 30 OFFSET %s',(batch_id,(page-1)*30))
        for f in row['filas']:
            f['propuesta']=json.loads(f['propuesta']); f['errores']=json.loads(f['errores'])
        return row


def pending(empresa,row_id):
    with s.db.transaction() as c:
        row=s.db.one(c,'''SELECT f.* FROM sh_import_filas f JOIN sh_importaciones i ON i.id=f.importacion_id
            WHERE f.id=%s AND i.empresa_id=%s''',(row_id,empresa))
        if not row: raise s.Error('Fila no encontrada.',404)
        row['original']=json.loads(row['original'])
        cats=s.db.all_rows(c,'SELECT * FROM sh_catalogos WHERE empresa_id=%s',(empresa,))
        externals=s.external_people(c,empresa)
        row['propuesta'],row['errores']=suggest(row['original'],s.employees(c,empresa,False),cats,externals)
        row['propuesta']['envio_id']=source_id(empresa,row['original'])
        return row


def import_ready(empresa,batch_id,user):
    with s.db.transaction() as c:
        if not s.db.one(c,'SELECT id FROM sh_importaciones WHERE id=%s AND empresa_id=%s',(batch_id,empresa)):
            raise s.Error('Importación no encontrada.',404)
        rows=s.db.all_rows(c,'SELECT id,original FROM sh_import_filas WHERE importacion_id=%s AND evento_id IS NULL',(batch_id,))
        people=s.employees(c,empresa,False)
        cats=s.db.all_rows(c,'SELECT * FROM sh_catalogos WHERE empresa_id=%s',(empresa,))
        externals=s.external_people(c,empresa)
    ok=0; unresolved=0
    failed=[]
    for r in rows:
        original=json.loads(r['original'])
        proposed,errors=suggest(original,people,cats,externals)
        proposed['envio_id']=source_id(empresa,original)
        if errors:
            unresolved+=1; failed.append((s.dumps(errors),r['id'])); continue
        try:
            _,created=s.create(empresa,proposed,user=user,import_row=r['id'])
            ok+=int(created)
        except s.Error as exc:
            unresolved+=1
            failed.append((s.dumps([str(exc)]),r['id']))
    if failed:
        with s.db.transaction() as c:
            c.executemany('UPDATE sh_import_filas SET errores=%s WHERE id=%s',failed)
    return ok,unresolved
