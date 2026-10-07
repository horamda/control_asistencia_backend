"""Plantilla editable e histórico corregible; no modifica el archivo ni la base original."""
import io
import json
import re
from datetime import date
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

from services import seguridad_service as s
from services import seguridad_import_service as importer
from services.seguridad_excel_schema import HEADERS


def literal(sheet, row):
    """Nunca ejecutar como fórmula texto proveniente de empleados o respuestas."""
    sheet.append(row)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def build(empresa, people, catalogs, records=(), externals=()):
    book = Workbook()
    instructions = book.active
    instructions.title = 'Instrucciones'
    carga = book.create_sheet('Carga')
    staff = book.create_sheet('Empleados')
    categories = book.create_sheet('Categorías')
    external_sheet = book.create_sheet('Externos')
    lists = book.create_sheet('Listas')
    original_sheet = book.create_sheet('Original') if records else None
    instructions.append(['IMPORTACIÓN DE SEGURIDAD E HIGIENE'])
    guidance = [
        'Complete la hoja Carga. No cambie encabezados, ID registro ni Fila origen.',
        'Una fila representa un evento. Para varias personas, separe los legajos con punto y coma: 0010;0011.',
        'Legajos se guarda como texto: conserve los ceros iniciales. Consulte Empleados; no se asignan por nombre aproximado.',
        'Tipo persona: empleado, externo o mixto. Pendiente significa que todavía debe identificar a la persona.',
        'Para externos nuevos complete Nombre externo. Empresa externa es opcional; no se exige legajo.',
        'Si el externo ya está registrado, use su ID de la hoja Externos; separe varios IDs con punto y coma y deje Nombre externo vacío.',
        'No marque como externo a alguien solo porque no encuentra su legajo: podría ser un exempleado o un nombre escrito distinto.',
        'Tipo: seguro, inseguro, accidente o incidente. Fecha y descripción son obligatorias para todos.',
        'Seguro/inseguro: categoría del mismo tipo. Inseguro: Advertido si/no. Accidente/incidente: hora y lugar.',
        'Fechas: use una fecha de Excel o AAAA-MM-DD. Hora: HH:MM. No se admiten fechas futuras.',
        'Revisión es una ayuda que se recalcula en Excel; la validación definitiva se realiza al subir.',
        'Si hay errores, la plantilla se rechaza completa y el panel muestra fila y campo. No se crean eventos parciales.',
        'Puede retirar temporalmente de Carga las filas incompletas. El histórico original permanece en su lote y en Original.',
        'Los eventos ya importados no se duplican si conserva sus IDs y datos. Para cambiar un evento existente, use el panel.',
        'Las categorías y empleados reflejan la fecha de descarga. Descargue otra plantilla si cambian los catálogos.',
        'Las listas de legajos incluyen empleados inactivos para poder recuperar el histórico.',
        'Enlaces originales conserva referencias a fotos; no adjunta ni descarga imágenes privadas.',
        'Después de subir, revise el lote y cree los eventos. Todos quedan pendientes de aprobación.',
        'Para registros nuevos use los IDs de una plantilla nueva; no copie IDs de otros eventos.',
    ]
    for text in guidance:
        instructions.append([text])
    instructions.column_dimensions['A'].width = 115
    instructions.row_dimensions[1].height = 30
    for row in range(2, len(guidance)+2):
        instructions.row_dimensions[row].height = 34
        instructions.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
    staff.append(['Legajo', 'Apellido y nombre', 'Sucursal', 'Sector', 'Puesto', 'Activo'])
    for person in people:
        literal(staff, [str(person['legajo'] or ''), f"{person['apellido']} {person['nombre']}",
                        person.get('sucursal_nombre'), person.get('sector_nombre'), person.get('puesto_nombre'),
                        'Sí' if person['activo'] else 'No'])
    categories.append(['ID', 'Tipo', 'Categoría'])
    active = [c for c in catalogs if c['activo'] and c['clase'] == 'categoria']
    for cat in active:
        literal(categories, [cat['id'], cat['tipo'], cat['nombre']])
    external_sheet.append(['ID externo','Nombre','Empresa / procedencia','Activo'])
    for person in externals:
        literal(external_sheet,[person['id'],person['nombre'],person['empresa'],'Sí' if person['activo'] else 'No'])
    lists.append(['Tipos', 'Advertido', 'seguro', 'inseguro', 'accidente', 'incidente'])
    kind_order = ['seguro','inseguro','accidente','incidente']
    list_columns = [kind_order, ['si', 'no']] + [[c['nombre'] for c in active if c['tipo']==tipo] for tipo in kind_order]
    for idx in range(max([len(v) for v in list_columns] + [1])):
        literal(lists, [col[idx] if idx < len(col) else None for col in list_columns])
    for column, name in enumerate(['Tipos', 'Avisos', 'seguro', 'inseguro', 'accidente', 'incidente'], 1):
        letter = get_column_letter(column)
        book.defined_names.add(DefinedName(name, attr_text=f"'Listas'!${letter}$2:${letter}${max(2,len(list_columns[column-1])+1)}"))
    book.defined_names.add(DefinedName('Legajos', attr_text=f"'Empleados'!$A$2:$A${max(2,len(people)+1)}"))
    carga.append(HEADERS)
    persons = {p['id']: p for p in people}
    cats = {c['id']: c for c in catalogs}
    originals = [json.loads(r['original']) if isinstance(r['original'],str) else r['original'] for r in records]
    common_headers = originals[0]['encabezados'] if originals else []
    same_headers = all(r['encabezados']==common_headers for r in originals)
    if original_sheet:
        literal(original_sheet,['Fila origen', 'Hoja', 'Fila Excel'] + (common_headers if same_headers else ['Campo original','Valor original']))
    for record in records:
        raw = record['original']
        if isinstance(raw, str):
            raw = json.loads(raw)
        proposed, errors = importer.suggest(raw, people, catalogs,externals)
        external_people=proposed.get('externos',[])
        new_external=next((p for p in external_people if not p.get('id')), {})
        person_type='mixto' if external_people and proposed['involucrados'] else 'externo' if external_people else 'empleado' if proposed['involucrados'] else 'pendiente'
        fecha = proposed.get('fecha_evento') or ''
        try:
            fecha = date.fromisoformat(fecha)
        except ValueError:
            pass
        links = sorted(set(re.findall(r'https?://[^\s<>"\]]+', s.dumps(raw))))
        literal(carga, [importer.source_id(empresa, raw), record['id'],
                       ';'.join(str(persons[i]['legajo']) for i in proposed['involucrados']),
                       proposed.get('persona_original'), proposed.get('tipo'), fecha,
                       proposed.get('hora_evento'), cats.get(proposed.get('categoria_id'), {}).get('nombre', ''),
                       proposed.get('descripcion'), proposed.get('advertido'), proposed.get('lugar'),
                       proposed.get('clasificacion'), proposed.get('atencion'), proposed.get('lesion'),
                       proposed.get('parte_cuerpo'), '\n'.join(links),None,person_type,
                       ';'.join(str(p['id']) for p in external_people if p.get('id')),
                       new_external.get('nombre',''),new_external.get('empresa','')])
        if same_headers:
            literal(original_sheet, [record['id'],record['hoja'],record['fila']] + raw['celdas'])
        else:
            for index,value in enumerate(raw['celdas']):
                header=raw['encabezados'][index] if index<len(raw['encabezados']) else f'Columna {index+1}'
                literal(original_sheet,[record['id'],record['hoja'],record['fila'],header,value])
        carga.cell(carga.max_row,17).comment = Comment('Revisión al descargar: '+ ('; '.join(errors) or 'Datos completos') +
            (f". Ya vinculado al evento #{record['evento_id']}" if record.get('evento_id') else ''), 'FichaYa')
    if not records:
        for _ in range(100):
            literal(carga, [str(uuid4())])
            carga.cell(carga.max_row,18,'empleado')
    end = carga.max_row
    staff_end = max(2, staff.max_row)
    cats_end = max(2, categories.max_row)
    for row in range(2, end+1):
        checks = [
            f'IF(OR(R{row}="empleado",R{row}="mixto"),IF(C{row}="","Legajos; ",IF(ISNUMBER(SEARCH(";",C{row})),"Verificar varios legajos al subir; ",IF(COUNTIF(\'Empleados\'!$A$2:$A${staff_end},C{row})<>1,"Legajo inexistente/ambiguo; ",""))),"")',
            f'IF(OR(R{row}="",R{row}="pendiente"),"Tipo persona; ","")',
            f'IF(AND(OR(R{row}="externo",R{row}="mixto"),S{row}="",T{row}=""),"Nombre externo o ID; ","")',
            f'IF(AND(R{row}="externo",C{row}<>""),"Use mixto; ","")',
            f'IF(COUNTIF(Tipos,E{row})<>1,"Tipo; ","")',
            f'IFERROR(IF(OR(F{row}="",IF(ISNUMBER(F{row}),F{row},DATEVALUE(F{row}))<DATE(2000,1,1),IF(ISNUMBER(F{row}),F{row},DATEVALUE(F{row}))>TODAY()),"Fecha; ",""),"Fecha; ")',
            f'IF(OR(I{row}="",LEN(I{row})>8000),"Descripción; ","")',
            f'IF(AND(OR(E{row}="seguro",E{row}="inseguro"),COUNTIFS(\'Categorías\'!$B$2:$B${cats_end},E{row},\'Categorías\'!$C$2:$C${cats_end},H{row})<>1),"Categoría; ","")',
            f'IF(AND(E{row}="inseguro",COUNTIF(Avisos,J{row})<>1),"Advertido; ","")',
            f'IF(AND(OR(E{row}="accidente",E{row}="incidente"),OR(G{row}="",K{row}="")),"Hora/lugar; ","")',
            f'IF(OR(A{row}="",COUNTIF($A$2:$A${end},A{row})<>1),"ID registro; ","")',
        ]
        expr = '&'.join(checks)
        carga.cell(row,17, f'=IF(AND(COUNTA(C{row}:O{row},S{row}:U{row})=0,B{row}=""),"",IF(({expr})="","Lista para validar",{expr}))')
        for col in (1,3,7):
            carga.cell(row,col).number_format = '@'
        carga.cell(row,6).number_format = 'yyyy-mm-dd'
    for col, formula in [('C','=Legajos'),('E','=Tipos'),('H','=INDIRECT($E2)'),('J','=Avisos')]:
        validation = DataValidation(type='list', formula1=formula, allow_blank=True)
        validation.showErrorMessage = col != 'C'
        validation.errorTitle = 'Valor fuera de la lista'
        validation.error = 'Seleccione una opción válida. La categoría debe corresponder al tipo.'
        validation.showInputMessage = True
        validation.promptTitle = 'Legajo' if col=='C' else 'Seleccione de la lista'
        validation.prompt = 'Para varios legajos, sepárelos con ;. Conserve ceros iniciales.' if col=='C' else 'Complete los campos obligatorios antes de importar.'
        carga.add_data_validation(validation)
        validation.add(f'{col}2:{col}{end}')
    dates = DataValidation(type='date',operator='between',formula1='DATE(2000,1,1)',formula2='TODAY()',allow_blank=True)
    dates.showErrorMessage=True; dates.error='Use una fecha entre 2000 y hoy.'
    carga.add_data_validation(dates); dates.add(f'F2:F{end}')
    carga.conditional_formatting.add(f'Q2:Q{end}',FormulaRule(formula=['AND($Q2<>"",$Q2<>"Lista para validar")'],fill=PatternFill('solid',fgColor='FFF0CF')))
    carga.conditional_formatting.add(f'Q2:Q{end}',FormulaRule(formula=['$Q2="Lista para validar"'],fill=PatternFill('solid',fgColor='DFF2E6')))
    person_types=DataValidation(type='list',formula1='"empleado,externo,mixto,pendiente"',allow_blank=False)
    person_types.showErrorMessage=True
    carga.add_data_validation(person_types); person_types.add(f'R2:R{end}')
    widths = [39,13,18,30,16,15,12,36,65,13,28,24,24,24,24,40,55,18,24,35,30]
    for col,width in enumerate(widths,1):
        carga.column_dimensions[get_column_letter(col)].width=width
    for sheet in book:
        sheet.sheet_view.showGridLines=False
        sheet.freeze_panes='D2' if sheet==carga else 'A2'
        if sheet != instructions:
            sheet.auto_filter.ref=sheet.dimensions
            sheet.row_dimensions[1].height=30
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='17344D')
            cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        if sheet not in (instructions,carga):
            for column in range(1,sheet.max_column+1):
                sheet.column_dimensions[get_column_letter(column)].width=30 if column<4 else 60
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            if sheet != instructions:
                sheet.row_dimensions[row[0].row].height=45
    # Give long observations room without making the historical sheet unusable.
    for row in carga.iter_rows(min_row=2):
        lines=max(3, *(sum(max(1,(len(part)+widths[idx]-1)//widths[idx]) for part in str(row[idx].value or '').splitlines())
                       for idx in (3,7,8,10)))
        carga.row_dimensions[row[0].row].height=min(150,lines*15)
    carga['A1'].comment=Comment('Identificador estable. No editar ni copiar entre eventos. Evita duplicados al reimportar.', 'FichaYa')
    carga['B1'].comment=Comment('Vínculo con la respuesta histórica. No editar. Vacío en registros nuevos.', 'FichaYa')
    carga['T1'].comment=Comment('Único dato obligatorio para una persona externa nueva. Empresa externa es opcional. Si ya existe, use su ID y deje este campo vacío.', 'FichaYa')
    lists.sheet_state='hidden'
    book.active=0
    buffer=io.BytesIO(); book.save(buffer)
    return buffer.getvalue()


def export(empresa, batch_id=None):
    with s.db.transaction() as c:
        people=s.employees(c,empresa,False)
        catalogs=s.db.all_rows(c,'SELECT * FROM sh_catalogos WHERE empresa_id=%s ORDER BY clase,tipo,nombre',(empresa,))
        externals=s.external_people(c,empresa)
        records=[]
        if batch_id:
            batch=s.db.one(c,'SELECT id FROM sh_importaciones WHERE id=%s AND empresa_id=%s',(batch_id,empresa))
            if not batch:
                raise s.Error('Importación no encontrada.',404)
            records=s.db.all_rows(c,'SELECT * FROM sh_import_filas WHERE importacion_id=%s ORDER BY fila,id',(batch_id,))
    return build(empresa,people,catalogs,records,externals)
