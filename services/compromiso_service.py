"""Indice de compromiso (engagement) por empresa, año, sucursal y sector.

Cada fila es un porcentaje (0 a 100). sector_id = 0 es el total de una sucursal y sucursal_id = 0 (con sector 0) el de
la empresa: si no se cargan, se calculan como el promedio simple de los sectores. Se carga desde el panel (planilla
Excel o CSV, o a mano) y cada empleado ve en la app el de su sector, su sucursal y su empresa, con el año anterior."""
import csv
import datetime as dt
import io
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from repositories import carga_repository as db
from services.carga_service import now_local

ANIO_MIN, ANIO_MAX = 2000, 2100
MAX_FILAS_IMPORTACION = 2000
COLUMNAS = ("Año", "Sucursal", "Área", "Índice")
_ALIAS = {"ano": "anio", "anio": "anio", "year": "anio", "sucursal": "sucursal", "area": "area", "sector": "area",
          "indice": "indice", "indice de compromiso": "indice", "compromiso": "indice", "engagement": "indice",
          "valor": "indice"}


class CompromisoError(ValueError):
    pass


def _norm(texto):
    sin_acentos = "".join(c for c in unicodedata.normalize("NFKD", str(texto or "")) if not unicodedata.combining(c))
    return " ".join(sin_acentos.lower().split())


def _anio(valor):
    try:
        anio = int(float(str(valor).strip()))
    except (TypeError, ValueError):
        raise CompromisoError("Año inválido.") from None
    if not ANIO_MIN <= anio <= ANIO_MAX:
        raise CompromisoError(f"Año fuera de rango ({ANIO_MIN} a {ANIO_MAX}).")
    return anio


def _indice(valor):
    texto = str(valor if valor is not None else "").strip().replace("%", "").replace(",", ".")
    try:
        indice = float(texto)
    except ValueError:
        raise CompromisoError("Índice inválido: use un número de 0 a 100.") from None
    if 0 < indice < 1 and "." in texto and not str(valor).strip().endswith("%"):
        indice *= 100  # 0,72 en la planilla = 72 %
    if not 0 <= indice <= 100:
        raise CompromisoError("El índice tiene que estar entre 0 y 100.")
    return round(indice, 2)


def catalogo(empresa_id):
    """Sucursales y sectores activos de la empresa (para formularios, importacion y plantilla)."""
    with db.transaction(read_only=True) as c:
        sucursales = db.all_rows(c, "SELECT id, nombre FROM sucursales WHERE empresa_id = %s AND activa = 1 ORDER BY nombre",
                                 (int(empresa_id),))
        sectores = db.all_rows(c, "SELECT id, nombre FROM sectores WHERE empresa_id = %s AND activo = 1 ORDER BY nombre",
                               (int(empresa_id),))
    return sucursales, sectores


def anios(empresa_id):
    with db.transaction(read_only=True) as c:
        return [int(r["anio"]) for r in db.all_rows(
            c, "SELECT DISTINCT anio FROM compromiso_indices WHERE empresa_id = %s ORDER BY anio DESC", (int(empresa_id),))]


def _filas(c, empresa_id, anio):
    return db.all_rows(c, """SELECT k.id, k.anio, k.sucursal_id, k.sector_id, k.indice, k.actualizado_at,
                                    su.nombre AS sucursal, se.nombre AS sector
                             FROM compromiso_indices k
                             LEFT JOIN sucursales su ON su.id = k.sucursal_id
                             LEFT JOIN sectores se ON se.id = k.sector_id
                             WHERE k.empresa_id = %s AND k.anio = %s
                             ORDER BY k.sucursal_id = 0 DESC, su.nombre, k.sector_id = 0 DESC, se.nombre""",
                       (int(empresa_id), int(anio)))


def _promedio(valores):
    valores = [float(v) for v in valores]
    return round(sum(valores) / len(valores), 1) if valores else None


def _resumen(filas):
    """Totales por sucursal y de la empresa: el cargado, o el promedio de sus sectores."""
    areas = [f for f in filas if f["sucursal_id"] and f["sector_id"]]
    cargados = {(f["sucursal_id"], f["sector_id"]): float(f["indice"]) for f in filas}
    sucursales = {}
    for f in filas:
        if f["sucursal_id"]:
            sucursales.setdefault(f["sucursal_id"], f["sucursal"] or f"Sucursal #{f['sucursal_id']}")
    por_sucursal = []
    for sid, nombre in sorted(sucursales.items(), key=lambda x: x[1].casefold()):
        propias = [f for f in areas if f["sucursal_id"] == sid]
        total = cargados.get((sid, 0))
        por_sucursal.append(dict(sucursal_id=sid, nombre=nombre, sectores=propias,
                                 indice=round(total, 1) if total is not None else _promedio(f["indice"] for f in propias),
                                 calculado=total is None))
    total = cargados.get((0, 0))
    empresa = dict(indice=round(total, 1) if total is not None else _promedio(f["indice"] for f in areas),
                   calculado=total is None)
    return dict(empresa=empresa, sucursales=por_sucursal)


def listado(empresa_id, anio):
    with db.transaction(read_only=True) as c:
        filas = _filas(c, empresa_id, anio)
    return dict(filas=filas, **_resumen(filas))


def _validar_lugar(c, empresa_id, sucursal_id, sector_id):
    sucursal_id, sector_id = int(sucursal_id or 0), int(sector_id or 0)
    if not sucursal_id and sector_id:
        raise CompromisoError("Para cargar un área elegí también la sucursal.")
    if sucursal_id and not db.one(c, "SELECT id FROM sucursales WHERE id = %s AND empresa_id = %s", (sucursal_id, int(empresa_id))):
        raise CompromisoError("La sucursal no es de esta empresa.")
    if sector_id and not db.one(c, "SELECT id FROM sectores WHERE id = %s AND empresa_id = %s", (sector_id, int(empresa_id))):
        raise CompromisoError("El área no es de esta empresa.")
    return sucursal_id, sector_id


def _guardar(c, empresa_id, user_id, anio, sucursal_id, sector_id, indice):
    c.execute("""INSERT INTO compromiso_indices (empresa_id, anio, sucursal_id, sector_id, indice, actualizado_por, actualizado_at)
                 VALUES (%s, %s, %s, %s, %s, %s, %s)
                 ON DUPLICATE KEY UPDATE indice = VALUES(indice), actualizado_por = VALUES(actualizado_por),
                                         actualizado_at = VALUES(actualizado_at)""",
              (int(empresa_id), anio, sucursal_id, sector_id, indice, user_id, now_local().replace(microsecond=0)))


def guardar(empresa_id, user_id, anio, sucursal_id, sector_id, indice):
    """Carga o reemplaza una fila (la misma empresa, año, sucursal y area se actualiza)."""
    anio, indice = _anio(anio), _indice(indice)
    with db.transaction() as c:
        sucursal_id, sector_id = _validar_lugar(c, empresa_id, sucursal_id, sector_id)
        _guardar(c, empresa_id, user_id, anio, sucursal_id, sector_id, indice)
    return anio


def borrar(empresa_id, fila_id):
    with db.transaction() as c:
        c.execute("DELETE FROM compromiso_indices WHERE id = %s AND empresa_id = %s", (int(fila_id), int(empresa_id)))
        if not c.rowcount:
            raise CompromisoError("Fila no encontrada.")


def _leer_filas(nombre, contenido):
    if (nombre or "").lower().endswith(".xlsx"):
        try:
            hoja = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True).active
        except Exception:
            raise CompromisoError("No se pudo leer la planilla: use un archivo .xlsx o .csv.") from None
        filas = [list(r) for r in hoja.iter_rows(values_only=True)]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        separador = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=separador))
    if not filas:
        raise CompromisoError("El archivo está vacío.")
    encabezado = [_ALIAS.get(_norm(h)) for h in filas[0]]
    faltan = {"anio", "sucursal", "indice"} - set(encabezado)
    if faltan:
        raise CompromisoError("Faltan columnas: " + ", ".join(sorted({"anio": "Año", "sucursal": "Sucursal", "indice": "Índice"}[f] for f in faltan))
                              + ". La planilla lleva Año, Sucursal, Área e Índice.")
    datos = []
    for numero, valores in enumerate(filas[1:], start=2):
        fila = {k: v for k, v in zip(encabezado, valores) if k}
        if all(v in (None, "") for v in fila.values()):
            continue
        datos.append((numero, fila))
    if len(datos) > MAX_FILAS_IMPORTACION:
        raise CompromisoError(f"Máximo {MAX_FILAS_IMPORTACION} filas por archivo.")
    return datos


def importar(empresa_id, user_id, nombre, contenido):
    """Importa la planilla. Si alguna fila tiene un error no se guarda nada y se devuelven todos los errores.
    Sucursal vacia o "Toda la empresa" = total de la empresa; Area vacia o "Toda la sucursal" = total de la sucursal."""
    datos = _leer_filas(nombre, contenido)
    sucursales, sectores = catalogo(empresa_id)
    por_nombre_suc = {_norm(s["nombre"]): s["id"] for s in sucursales}
    por_nombre_sec = {_norm(s["nombre"]): s["id"] for s in sectores}
    validas, errores = [], []
    for numero, fila in datos:
        try:
            anio = _anio(fila.get("anio"))
            indice = _indice(fila.get("indice"))
            suc, area = _norm(fila.get("sucursal")), _norm(fila.get("area"))
            if suc in ("", "toda la empresa", "empresa", "total"):
                if area not in ("", "toda la sucursal", "total"):
                    raise CompromisoError("un área necesita su sucursal")
                sucursal_id = sector_id = 0
            else:
                if suc not in por_nombre_suc:
                    raise CompromisoError(f"sucursal desconocida «{fila.get('sucursal')}»")
                sucursal_id = por_nombre_suc[suc]
                if area in ("", "toda la sucursal", "total"):
                    sector_id = 0
                elif area in por_nombre_sec:
                    sector_id = por_nombre_sec[area]
                else:
                    raise CompromisoError(f"área desconocida «{fila.get('area')}»")
            validas.append((anio, sucursal_id, sector_id, indice))
        except CompromisoError as exc:
            errores.append(f"Fila {numero}: {str(exc)[0].lower() + str(exc)[1:]}")
    if errores:
        raise CompromisoError("No se importó nada. " + " · ".join(errores[:20]) + (" · …" if len(errores) > 20 else ""))
    if not validas:
        raise CompromisoError("La planilla no tiene filas con datos.")
    with db.transaction() as c:
        for anio, sucursal_id, sector_id, indice in validas:
            _guardar(c, empresa_id, user_id, anio, sucursal_id, sector_id, indice)
    return dict(filas=len(validas), anios=sorted({v[0] for v in validas}))


def plantilla(empresa_id, anio=None):
    """Planilla con las columnas y una fila por sucursal y area con empleados activos, lista para completar."""
    anio = anio or now_local().year
    with db.transaction(read_only=True) as c:
        combinaciones = db.all_rows(c, """SELECT DISTINCT su.nombre AS sucursal, se.nombre AS sector
                                          FROM empleados e JOIN sucursales su ON su.id = e.sucursal_id
                                          JOIN sectores se ON se.id = e.sector_id
                                          WHERE e.empresa_id = %s AND e.activo = 1 ORDER BY su.nombre, se.nombre""",
                                    (int(empresa_id),))
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Compromiso"
    hoja.append(list(COLUMNAS))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="0866FF")
    for fila in combinaciones:
        hoja.append([anio, fila["sucursal"], fila["sector"], None])
    for col, ancho in zip("ABCD", (8, 28, 28, 10)):
        hoja.column_dimensions[col].width = ancho
    nota = libro.create_sheet("Cómo completar")
    for linea in ("Índice: número de 0 a 100 (también se acepta 0,72 o 72%).",
                  "Área vacía = total de la sucursal. Sucursal vacía = total de la empresa.",
                  "Si no se carga un total, se calcula como el promedio de las áreas.",
                  "Una fila con la misma sucursal, área y año reemplaza a la anterior."):
        nota.append([linea])
    nota.column_dimensions["A"].width = 90
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def para_empleado(empleado, anio=None):
    """Lo que ve el empleado en la app: su area (sector en su sucursal), su sucursal y la empresa, con el año
    anterior. Sin año, el ultimo con datos de su empresa."""
    empresa_id = int(empleado["empresa_id"])
    disponibles = anios(empresa_id)
    anio = int(anio) if anio else (disponibles[0] if disponibles else now_local().year)
    sucursal_id, sector_id = int(empleado.get("sucursal_id") or 0), int(empleado.get("sector_id") or 0)
    with db.transaction(read_only=True) as c:
        nombres = db.one(c, """SELECT su.nombre AS sucursal, se.nombre AS sector, emp.razon_social AS empresa
                               FROM empleados e LEFT JOIN sucursales su ON su.id = e.sucursal_id
                               LEFT JOIN sectores se ON se.id = e.sector_id LEFT JOIN empresas emp ON emp.id = e.empresa_id
                               WHERE e.id = %s""", (int(empleado["id"]),)) or {}
        actual, anterior = _filas(c, empresa_id, anio), _filas(c, empresa_id, anio - 1)

    def niveles(filas):
        resumen = _resumen(filas)
        area = next((f for f in filas if f["sucursal_id"] == sucursal_id and f["sector_id"] == sector_id and sector_id), None)
        suc = next((s for s in resumen["sucursales"] if s["sucursal_id"] == sucursal_id), None)
        return dict(area=dict(indice=round(float(area["indice"]), 1) if area else None, calculado=False),
                    sucursal=dict(indice=suc["indice"], calculado=suc["calculado"]) if suc else dict(indice=None, calculado=False),
                    empresa=resumen["empresa"])

    hoy, antes = niveles(actual), niveles(anterior)

    def nivel(clave, nombre):
        valor, previo = hoy[clave]["indice"], antes[clave]["indice"]
        return dict(nombre=nombre, indice=valor, calculado=hoy[clave]["calculado"], anterior=previo,
                    variacion=round(valor - previo, 1) if valor is not None and previo is not None else None)

    return dict(anio=anio, anios_disponibles=disponibles,
                area=nivel("area", nombres.get("sector") or "Sin área"),
                sucursal=nivel("sucursal", nombres.get("sucursal") or "Sin sucursal"),
                empresa=nivel("empresa", nombres.get("empresa") or "Empresa"))
